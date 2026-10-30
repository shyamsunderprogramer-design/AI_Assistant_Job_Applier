"""Read a resume or job description file into ordered blocks of text, each with its source.

Every block says where it came from -- "paragraph 12", "table 1, row 2, cell 1",
"page 2, line 7", "header" -- so a match or a generated sentence can point at
the part of the resume that supports it.

DOCX is read in document order (paragraphs and tables interleaved as they
appear), with list items, numbering, headings, bold, font size, hyperlinks,
headers and footers. PDF text is read page by page; a page with no selectable
text is read with macOS text recognition when it is available, and every
block that came from recognition is marked low confidence.

Every failure is a ReadError with a kind the page can act on: missing,
unsupported, damaged, empty or unreadable.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = (".docx", ".pdf", ".txt", ".md")
OCR_MIN_CHARS = 40          # a page with less selectable text than this is treated as scanned


@dataclass
class Block:
    text: str
    source: str
    style: str = ""
    heading: bool = False
    list_item: bool = False
    numbered: bool = False
    bold: bool = False
    size: float | None = None
    links: list[str] = field(default_factory=list)
    confidence: str = "high"            # "low" when it came from text recognition


@dataclass
class Extracted:
    kind: str                           # docx | pdf | text
    blocks: list[Block]
    warnings: list[str] = field(default_factory=list)
    ocr: bool = False

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


class ReadError(Exception):
    """A file that cannot be read, with a kind and a message for the person."""

    def __init__(self, kind: str, message: str):
        super().__init__(message)
        self.kind = kind


def read_document(path: Path | str, name: str | None = None) -> Extracted:
    path = Path(path)
    label = name or path.name
    if not path.exists():
        raise ReadError("missing", f"{label}: the file was not found. Upload it again.")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ReadError("unsupported", f"{label}: {suffix or 'this'} files are not supported. "
                                       f"Use a .docx or .pdf.")
    if path.stat().st_size == 0:
        raise ReadError("empty", f"{label}: the file is empty.")
    if suffix == ".docx":
        result = _read_docx(path, label)
    elif suffix == ".pdf":
        result = _read_pdf(path, label)
    else:
        lines = [l.strip() for l in path.read_text(encoding="utf-8", errors="replace").splitlines()]
        result = Extracted("text", [Block(l, f"line {i}") for i, l in enumerate(lines, 1) if l])
    if not any(b.text.strip() for b in result.blocks):
        raise ReadError("empty", f"{label}: no text could be read from this file.")
    result.warnings += scramble_warnings(result.blocks)
    return result


# -- DOCX ------------------------------------------------------------------------

def _read_docx(path: Path, label: str) -> Extracted:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(str(path))
    except Exception as exc:
        raise ReadError("damaged", f"{label}: this Word file could not be opened — it may be "
                                   f"damaged or not really a .docx ({type(exc).__name__}).") from exc

    formats = _numbering_formats(document)
    blocks: list[Block] = []
    for section in document.sections[:1]:
        for p in section.header.paragraphs:
            b = _paragraph_block(p, "header", formats)
            if b:
                blocks.append(b)

    para_no = table_no = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para_no += 1
            b = _paragraph_block(Paragraph(child, document), f"paragraph {para_no}", formats)
            if b:
                blocks.append(b)
        elif tag == "tbl":
            table_no += 1
            table = Table(child, document)
            for r, row in enumerate(table.rows, 1):
                seen = set()
                for c, cell in enumerate(row.cells, 1):
                    if id(cell._tc) in seen:          # a merged cell repeats across columns
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        b = _paragraph_block(p, f"table {table_no}, row {r}, cell {c}", formats)
                        if b:
                            blocks.append(b)

    for section in document.sections[:1]:
        for p in section.footer.paragraphs:
            b = _paragraph_block(p, "footer", formats)
            if b:
                blocks.append(b)
    return Extracted("docx", blocks)


def _numbering_formats(document) -> dict[tuple[str, str], str]:
    """(numId, level) -> numFmt ("bullet", "decimal", ...), so a numbered list is told apart."""
    formats: dict[tuple[str, str], str] = {}
    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return formats
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    w = "{%s}" % ns["w"]
    abstract = {}
    for a in numbering.findall("w:abstractNum", ns):
        levels = {lvl.get(w + "ilvl"): (lvl.find("w:numFmt", ns).get(w + "val")
                                        if lvl.find("w:numFmt", ns) is not None else "")
                  for lvl in a.findall("w:lvl", ns)}
        abstract[a.get(w + "abstractNumId")] = levels
    for num in numbering.findall("w:num", ns):
        ref = num.find("w:abstractNumId", ns)
        if ref is None:
            continue
        for level, fmt in abstract.get(ref.get(w + "val"), {}).items():
            formats[(num.get(w + "numId"), level)] = fmt
    return formats


def _paragraph_block(p, source: str, formats) -> Block | None:
    text = re.sub(r"[ \t]+", " ", p.text or "").strip()
    if not text:
        return None
    style = (p.style.name if p.style is not None else "") or ""
    numbered = list_item = False
    num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
    if num_pr is not None and num_pr.numId is not None:
        list_item = True
        level = num_pr.ilvl.val if num_pr.ilvl is not None else 0
        fmt = formats.get((str(num_pr.numId.val), str(level)), "bullet")
        numbered = fmt not in ("bullet", "none", "")
    elif "list" in style.lower():
        list_item = True
        numbered = "number" in style.lower()
    runs = [r for r in p.runs if r.text.strip()]
    bold = bool(runs) and all(bool(r.bold) or bool(p.style is not None and p.style.font.bold)
                              for r in runs)
    size = next((r.font.size.pt for r in runs if r.font.size is not None), None)
    if size is None and p.style is not None and p.style.font.size is not None:
        size = p.style.font.size.pt
    links = []
    for link in getattr(p, "hyperlinks", []) or []:
        url = getattr(link, "url", "") or ""
        if url:
            links.append(url)
    return Block(text=text, source=source, style=style,
                 heading=style.lower().startswith(("heading", "title")),
                 list_item=list_item, numbered=numbered, bold=bold, size=size, links=links)


# -- PDF -------------------------------------------------------------------------

def _pdfminer_page(path: Path, index: int) -> str | None:
    """One page's text by pdfminer, or None when it cannot be had.

    pypdf reads a resume's printed lines one by one and splits words at kerning gaps -- "Cloud
    Bui ld", "SonarQ ube", "scan ning" -- and a resume imported that way kept the breaks in every
    line. pdfminer, with words grouped generously, reads the same page whole.
    """
    try:
        from pdfminer.high_level import extract_text
        from pdfminer.layout import LAParams
        return extract_text(str(path), page_numbers=[index], laparams=LAParams(char_margin=4.0, word_margin=0.2))
    except Exception:
        return None


_BULLET = re.compile(r"^[•●▪◦‣∙·\-*]\s+")
_LABEL = re.compile(r"^[A-Z][\w/&.+()# -]{1,60}:\s")      # "Containerization & Orchestration: Docker…"
_ENDS_OPEN = re.compile(r"[,\-–—&/(]$|\b(and|or|of|the|to|for|with|in|on|by|a|an)$", re.I)


def _core(token: str) -> str:
    return token.strip(".,;:()[]\"'")


def mend_words(lines_text: str, whole: str) -> str:
    """pypdf's lines with its broken words put back the way pdfminer read them.

    pypdf keeps a resume's lines in their places -- a job title with its dates beside it -- but
    breaks words at wide letter gaps ("Bui ld", "experien ce.", "frameworks ,", "high -stakes",
    "20 –"). pdfminer reads the words whole but moves dates into a column of their own. Two
    pieces become one where pdfminer has the joined word and never the two side by side.
    """
    words = whole.split()
    known = {_core(w) for w in words} | set(words)
    pairs = {(a, b) for a, b in zip(words, words[1:])}
    out = []
    for line in lines_text.splitlines():
        tokens = line.split()
        merged: list[str] = []
        for token in tokens:
            if merged:
                joined = merged[-1] + token
                if (joined in known or _core(joined) in known) and (merged[-1], token) not in pairs \
                        and _core(joined):
                    merged[-1] = joined
                    continue
            merged.append(token)
        out.append(" ".join(merged))
    return "\n".join(out)


def logical_lines(text: str) -> list[str]:
    """A page's printed lines put back into the items they were written as.

    A PDF keeps no paragraphs, only lines, so one bullet wrapped over three lines came back as
    three bullets, each a fragment. A line continues the item above it when that item is a
    bullet and the line starts in lower case, or the item stops mid-phrase (a comma, a hyphen,
    "and"), or the item has not reached its full stop and the line is not a "Label: …" row.
    A line opening with a bullet, or a heading, job or date line, always starts afresh.
    """
    out: list[str] = []
    open_bullet = False
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            continue
        if _BULLET.match(line):
            out.append("• " + _BULLET.sub("", line))
            open_bullet = True
            continue
        prev = out[-1] if out else ""
        # A "Label: tools" row wrapped mid-list ("…, Helm," / "PodDisruptionBudgets, …") goes on too.
        if prev and not open_bullet and _LABEL.match(prev) and prev.endswith(",") and not _LABEL.match(line) \
                and not _BULLET.match(line):
            out[-1] = prev + " " + line
            continue
        joins = open_bullet and prev and (
            line[:1].islower() or _ENDS_OPEN.search(prev)
            or (not prev.endswith((".", "!", "?")) and not _LABEL.match(line) and not line.isupper()
                and not line.endswith(":")
                and not re.search(r"\b(19|20)\d\d\b", line)))
        if joins:
            glue = "" if re.search(r"\w[-–—]$", prev) else " "
            out[-1] = prev + glue + line
        else:
            out.append(line)
            open_bullet = False
    return out


def _read_pdf(path: Path, label: str) -> Extracted:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ReadError("unreadable", f"{label}: this PDF is password-protected.")
        pages = list(reader.pages)
    except ReadError:
        raise
    except Exception as exc:
        raise ReadError("damaged", f"{label}: this PDF could not be opened — it may be damaged "
                                   f"({type(exc).__name__}).") from exc
    if not pages:
        raise ReadError("empty", f"{label}: this PDF has no pages.")

    blocks: list[Block] = []
    warnings: list[str] = []
    scanned_pages = []
    for number, page in enumerate(pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        whole = _pdfminer_page(path, number - 1)
        if whole:
            text = mend_words(text, whole) if text.strip() else whole
        if len(text.strip()) < OCR_MIN_CHARS:
            scanned_pages.append(number)
            continue
        for line_no, line in enumerate(logical_lines(text), 1):
            blocks.append(Block(line, f"page {number}, line {line_no}"))

    ocr_used = False
    if scanned_pages:
        recognised = _ocr_pages(path, scanned_pages)
        if recognised is None:
            if not blocks:
                raise ReadError("unreadable", f"{label}: this PDF is a scanned image with no "
                                              f"selectable text, and text recognition is not "
                                              f"available on this computer. Upload a .docx or a "
                                              f"text PDF.")
            warnings.append(f"Pages {scanned_pages} are scanned images and could not be read.")
        else:
            ocr_used = True
            blocks += recognised
            warnings.append(f"Page(s) {', '.join(map(str, scanned_pages))} were scanned images; "
                            f"their text was recognised automatically and may contain errors. "
                            f"Check the extracted text below.")
            blocks.sort(key=lambda b: _page_of(b.source))
    return Extracted("pdf", blocks, warnings, ocr=ocr_used)


def _page_of(source: str) -> tuple[int, int]:
    m = re.match(r"page (\d+), line (\d+)", source)
    return (int(m.group(1)), int(m.group(2))) if m else (0, 0)


def _ocr_pages(path: Path, numbers: list[int]) -> list[Block] | None:
    """Recognise the text of scanned pages, top to bottom. None if recognition is unavailable."""
    try:
        import pypdfium2 as pdfium
        from ocrmac import ocrmac
    except Exception:
        return None
    blocks: list[Block] = []
    try:
        pdf = pdfium.PdfDocument(str(path))
        for number in numbers:
            image = pdf[number - 1].render(scale=2.5).to_pil()
            # The accurate model can come back empty while the fast one reads the page: on
            # 22 Oct 2026 it returned nothing even for plain 48pt Arial. Fast text beats none.
            found = (ocrmac.OCR(image, recognition_level="accurate").recognize()
                     or ocrmac.OCR(image, recognition_level="fast").recognize())
            # Vision boxes are (x, y, w, h) with the origin at the bottom-left.
            lines = _group_lines(found)
            for line_no, text in enumerate(lines, 1):
                blocks.append(Block(text, f"page {number}, line {line_no} (recognised)",
                                    confidence="low"))
    except Exception:
        return None
    return blocks


def _group_lines(found) -> list[str]:
    """Words and phrases recognised on a page, joined into lines in reading order."""
    items = sorted(((1 - box[1] - box[3], box[0], text) for text, _conf, box in found),
                   key=lambda t: (round(t[0], 2), t[1]))
    lines: list[list[tuple[float, float, str]]] = []
    for top, left, text in items:
        if lines and abs(lines[-1][0][0] - top) < 0.008:
            lines[-1].append((top, left, text))
        else:
            lines.append([(top, left, text)])
    return [" ".join(t for _, _, t in sorted(line, key=lambda i: i[1])).strip() for line in lines]


# -- uncertainty -------------------------------------------------------------------

def scramble_warnings(blocks: list[Block]) -> list[str]:
    """Text that came out garbled: broken words, stray symbols, letters run together."""
    words = re.findall(r"\S+", " ".join(b.text for b in blocks))
    if len(words) < 30:
        return []
    odd = [w for w in words if re.search(r"[^\w\s.,;:()/&+%#@'\-–—|•·$€£]", w)
           or (len(w) > 28 and "/" not in w and "." not in w)]
    singles = [w for w in words if len(w) == 1 and w.isalpha() and w.lower() not in ("a", "i")]
    if len(odd) / len(words) > 0.05 or len(singles) / len(words) > 0.15:
        return ["Some text looks scrambled (broken words or stray symbols). The PDF's text layer "
                "may be damaged — check the extracted text, or upload a .docx."]
    return []
