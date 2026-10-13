"""A resume health check: how well the resume is written and built, with no job posting.

The ATS score says how well a resume matches ONE posting. This says whether
the resume itself is sound, the things every posting's reader trips on:

  contact    an email, a phone, a city, a LinkedIn; nothing in the page header
             (some parsers skip headers); no full street address
  sections   Experience, Education and Skills under names a parser knows
             ("Work Experience", not "My Career Journey"); a summary
  impact     each bullet: starts with an action verb, not "Responsible for";
             shows a number; no "I" or "my"; not one line of three words or a
             paragraph of fifty; the same opening verb not used again and again;
             no "team player" or "results-driven"
  dates      every job dated, one date format, newest first, no dates in the
             future, gaps of more than six months pointed out
  length     pages against years of experience
  layout     (from the file) tables, text boxes, columns, images, unusual fonts,
             a PDF with no text in it
  polish     a doubled word, a doubled space, bullets that end with a full stop
             and bullets that do not, mixed in one resume

Each finding names the lines it is about, quoted, and how to fix it. The
score is out of 100: each category has a share and loses part of it per
finding. Nothing here changes the resume; a weak bullet can be rewritten with
the model on request (`rewrite_bullets`), from that bullet's own facts only.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

WEIGHTS = {"contact": 10, "sections": 15, "impact": 30, "dates": 10, "length": 5, "layout": 20, "polish": 10}
LABELS = {"contact": "Contact details", "sections": "Sections and headings", "impact": "Bullets and impact",
          "dates": "Dates", "length": "Length", "layout": "Layout an ATS can read", "polish": "Polish"}
COST = {"problem": 0.35, "warning": 0.15, "tip": 0.04}    # of a category's share, per finding

# A bullet that starts like this says what the job was, not what was done.
WEAK_STARTS = ("responsible for", "worked on", "worked with", "helped", "assisted", "involved in",
               "participated in", "tasked with", "duties included", "duties include", "in charge of",
               "was responsible", "were responsible", "handled", "did ", "was ", "were ", "had ",
               "exposure to", "familiar with", "knowledge of", "working on", "worked in")
PRONOUNS = re.compile(r"\b(I(?![/\-]O\b)|me|my|myself|we|our|us)\b")
# Phrases a recruiter has read ten thousand times; they show nothing.
CLICHES = ("team player", "hard working", "hard-working", "hardworking", "results-driven", "results driven",
           "detail-oriented", "detail oriented", "go-getter", "self-starter", "self starter",
           "think outside the box", "synergy", "dynamic individual", "proven track record",
           "excellent communication skills", "passionate about", "rockstar", "ninja", "guru",
           "fast learner", "quick learner", "strong work ethic", "best of breed", "go-to person")
VAGUE = ("various", "etc", "several", "many", "numerous", "a lot of", "lots of", "some")
ACTION_VERBS = set("""
accelerated achieved administered advanced advised analyzed analysed architected assembled assessed audited
authored automated balanced boosted built calculated captured chaired championed coached collaborated
completed composed conceived conducted configured consolidated constructed consulted containerized
converted coordinated created cut debugged decreased defined delivered deployed designed developed devised
diagnosed directed discovered documented doubled drafted drove eliminated enabled enforced engineered
enhanced established evaluated executed expanded expedited facilitated finalized fixed forecasted formed
founded generated grew guided halved hardened headed hired identified implemented improved increased
influenced initiated innovated inspected installed instituted integrated introduced invented investigated
launched led leveraged maintained managed mapped maximized mentored merged migrated minimized modeled
modelled modernized monitored motivated negotiated onboarded operated optimized orchestrated organized
overhauled oversaw partnered patched piloted pioneered planned prepared presented prevented prioritized
produced programmed promoted proposed prototyped provisioned published raised rebuilt recommended
reconciled recruited redesigned reduced refactored remediated reengineered reorganized replaced
reported researched resolved restructured revamped reviewed rewrote saved scaled scheduled scripted secured
served shipped simplified slashed solved spearheaded standardized steered streamlined strengthened
structured supervised supported surpassed taught tested tracked trained transformed translated tripled
troubleshot tuned unified upgraded validated verified won wrote
""".split())
# The same verbs in the present tense, for the current job: "Lead", "Build".
ACTION_VERBS |= {v[:-1] for v in ACTION_VERBS if v.endswith("ed") and v[-3] not in "ie"} | {
    "lead", "build", "run", "own", "drive", "design", "deploy", "manage", "architect", "automate",
    "maintain", "mentor", "implement", "develop", "migrate", "write", "oversee", "partner", "teach",
    "secure", "scale", "ship", "support", "test", "monitor", "operate", "deliver", "create"}
NUMBER = re.compile(r"\d|\b(two|three|four|five|six|seven|eight|nine|ten|twelve|dozens?|hundreds?|"
                    r"thousands?|millions?|billions?|half|double[sd]?|triple[sd]?)\b", re.I)
STANDARD = {"summary": "Summary", "experience": "Experience", "skills": "Skills", "education": "Education",
            "projects": "Projects", "certifications": "Certifications"}
COMMON_FONTS = {"arial", "calibri", "cambria", "georgia", "garamond", "helvetica", "times new roman",
                "verdana", "tahoma", "book antiqua", "palatino linotype", "trebuchet ms", "aptos",
                "segoe ui", "gill sans mt", "century gothic", "lato", "roboto", "open sans", "source sans pro",
                "dejavu sans", "liberation serif", "liberation sans", "carlito", "caladea"}
MONTH_WORD = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}", re.I)
MONTH_NUM = re.compile(r"\b\d{1,2}/\d{4}\b")
YEAR_ONLY = re.compile(r"(?<![/\w])\d{4}\s*(?:–|—|-|to)\s*(?:\d{4}|present|current|now)\b", re.I)


@dataclass
class Finding:
    category: str
    severity: str                  # problem | warning | tip
    title: str
    detail: str = ""
    fix: str = ""
    lines: list[str] = field(default_factory=list)


@dataclass
class Report:
    score: int
    categories: list[dict]
    findings: list[Finding]
    stats: dict

    def to_dict(self) -> dict:
        return {"score": self.score, "categories": self.categories,
                "findings": [asdict(f) for f in self.findings], "stats": self.stats}


def _clean(line: str) -> str:
    from ml.resume.parser import strip_bullet
    return strip_bullet(line or "").strip()


def _first_word(line: str) -> str:
    m = re.match(r"[A-Za-z][A-Za-z'-]*", _clean(line))
    return m.group(0).lower() if m else ""


# -- the parts of a resume the checks read -------------------------------------------

@dataclass
class Parts:
    """What the checks need, from a file (`from_file`) or from the builder's fields."""
    name: str = ""
    contact_text: str = ""                    # the lines above the first section
    headings: list[tuple[str, str | None]] = field(default_factory=list)   # (as written, kind or None)
    sections: dict = field(default_factory=dict)                            # kind -> [lines]
    bullets: list[str] = field(default_factory=list)                        # experience and project lines
    jobs: list[dict] = field(default_factory=list)                          # {header, start, end, present}
    text: str = ""
    pages: int | None = None
    header_footer: str = ""                   # text in a .docx page header or footer
    worked_years: float | None = None         # when known better than first-start-to-now (the builder)
    layout: list[Finding] = field(default_factory=list)


def from_file(path: Path | str) -> Parts:
    from ml.tailoring.reader import read_document
    from ml.tailoring.structure import build
    path = Path(path)
    extracted = read_document(path)
    model = build(extracted)
    parts = from_model(model, extracted)
    parts.layout = layout_findings(path, extracted)
    parts.header_footer = " ".join(b.text for b in extracted.blocks if b.source in ("header", "footer"))
    parts.pages = _pages(path)
    return parts


def from_model(model, extracted) -> Parts:
    from ml.tailoring.structure import _is_heading, section_kind
    from ml.resume.parser import looks_like_heading
    parts = Parts(name=model.name, text=extracted.text)
    seen_heading = False
    contact = []
    for i, block in enumerate(extracted.blocks):
        text = block.text.strip()
        if block.source in ("header", "footer"):
            continue
        kind = _is_heading(block)
        if kind:
            seen_heading = True
            parts.headings.append((text, kind))
            continue
        if not seen_heading:
            contact.append(text)
            continue
        # A heading-shaped line that names no section a parser knows: "My Career Journey".
        if (block.heading or (looks_like_heading(text) and text.isupper())) and not section_kind(text) \
                and len(text.split()) <= 5 and not re.search(r"\d", text) and not block.list_item:
            parts.headings.append((text, None))
    parts.contact_text = "\n".join(contact)
    for kind, ids in model.sections.items():
        parts.sections[kind] = [model.fact(fid).text for fid in ids if model.fact(fid)]
    for f in model.facts:
        marked = f.text.lstrip().startswith(("•", "-", "*", "▪", "◦", "‣"))
        if f.section in ("experience", "projects") and f.kind != "entry" and (
                f.kind == "bullet" or marked or (f.kind == "text" and len(f.text.split()) >= 8)):
            parts.bullets.append(f.text)
    parts.jobs = [{"header": e.header, "start": e.start, "end": e.end, "present": e.present,
                   "dates": e.dates} for e in model.entries if e.section == "experience"]
    return parts


def _pages(path: Path) -> int | None:
    try:
        if path.suffix.lower() == ".pdf":
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(str(path))
            try:
                return len(doc)
            finally:
                doc.close()
        if path.suffix.lower() == ".docx":
            from docx import Document
            doc = Document(str(path))
            words = sum(len(p.text.split()) for p in doc.paragraphs)
            return max(1, round(words / 550 + 0.4))
    except Exception:
        return None
    return None


# -- layout, from the file itself --------------------------------------------------------

def layout_findings(path: Path, extracted=None) -> list[Finding]:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return _docx_layout(path)
    if path.suffix.lower() == ".pdf":
        return _pdf_layout(path, extracted)
    return []


def _docx_layout(path: Path) -> list[Finding]:
    from docx import Document
    from docx.oxml.ns import qn
    out: list[Finding] = []
    doc = Document(str(path))
    body = doc.element.body
    if doc.tables:
        cells = sum(len(r.cells) for t in doc.tables for r in t.rows)
        out.append(Finding("layout", "problem", f"{len(doc.tables)} table(s) in the document",
                           f"{cells} cells. Many ATS read a table cell by cell, out of order, or not at all, "
                           "so a job title can land beside the wrong dates.",
                           "Put the content in ordinary lines; use a tab for dates on the right."))
    boxes = body.findall(".//" + qn("w:txbxContent"))
    if boxes:
        out.append(Finding("layout", "problem", f"{len(boxes)} text box(es)",
                           "Text in a text box is often skipped entirely by an ATS.",
                           "Move the text into the page as ordinary paragraphs."))
    pictures = body.findall(".//" + qn("w:drawing"))
    if pictures:
        out.append(Finding("layout", "warning", f"{len(pictures)} picture(s) or chart(s)",
                           "An ATS cannot read a picture; anything written in one (a logo, a skills chart, "
                           "a photo) is lost, and a photo invites bias in the US and UK.",
                           "Remove pictures; write skill levels as words."))
    for sect in body.findall(".//" + qn("w:sectPr")):
        cols = sect.find(qn("w:cols"))
        if cols is not None and int(cols.get(qn("w:num"), "1") or 1) > 1:
            out.append(Finding("layout", "problem", "The page is set in columns",
                               "A parser reads across the columns, mixing the left one's lines with the right's.",
                               "Use one column."))
            break
    hf = " ".join(p.text for s in doc.sections for p in list(s.header.paragraphs) + list(s.footer.paragraphs))
    if re.search(r"@|\d{3}[\s.)-]*\d{3}[\s.-]*\d{4}", hf):
        out.append(Finding("layout", "problem", "Contact details are in the page header or footer",
                           "Some ATS ignore the header and footer, so the resume arrives with no way to reach you.",
                           "Put the name and contact details in the body, at the top of the first page.",
                           [hf.strip()[:120]]))
    fonts = set()
    for el in body.iter(qn("w:rFonts")):
        name = el.get(qn("w:ascii")) or el.get(qn("w:hAnsi"))
        if name:
            fonts.add(name)
    try:
        default = doc.styles["Normal"].font.name
        if default:
            fonts.add(default)
    except KeyError:
        pass
    odd = sorted(f for f in fonts if f.lower() not in COMMON_FONTS and not f.lower().startswith(("+", "symbol")))
    if odd:
        out.append(Finding("layout", "tip", "Fonts a reader may not have: " + ", ".join(odd[:4]),
                           "A .docx shows a missing font as a substitute, which moves every line.",
                           "Use a common font (Calibri, Cambria, Arial, Georgia) or send the PDF."))
    if len(fonts) > 3:
        out.append(Finding("polish", "tip", f"{len(fonts)} different fonts",
                           "More than two fonts reads as pasted together.", "Use one font, or one for headings."))
    return out


def _pdf_layout(path: Path, extracted=None) -> list[Finding]:
    from ml.tailoring.documents import _lines
    out: list[Finding] = []
    try:
        lines = _lines(path)
    except Exception as exc:
        return [Finding("layout", "problem", "The PDF could not be read",
                        f"{type(exc).__name__}. An ATS will likely fail on it too.", "Export it again from Word.")]
    pages = {p for p, *_ in lines}
    chars = {p: 0 for p in pages}
    for p, *_, text in lines:
        chars[p] += len(text)
    empty = sorted(p for p, n in chars.items() if n < 40)
    if empty:
        out.append(Finding("layout", "problem", "Page(s) with no text an ATS can read: "
                           + ", ".join(map(str, empty)),
                           "These pages are pictures of text (a scan or a design export). An ATS sees a blank page.",
                           "Export the PDF from Word or Google Docs, not as an image."))
    # Columns: many lines of real text starting past the middle of the page.
    for page in sorted(pages):
        rows = [(l, r, w, t) for p, l, b, r, tp, w, t in lines if p == page and t]
        if not rows:
            continue
        width = rows[0][2]
        right = [t for l, r, w, t in rows if l > width * 0.45 and len(t) > 25]
        left = [t for l, r, w, t in rows if l < width * 0.2 and r < width * 0.55 and len(t) > 25]
        if len(right) >= 5 and len(left) >= 5:
            out.append(Finding("layout", "problem", f"Page {page} looks like two columns",
                               "A parser reads straight across both columns and mixes them.",
                               "Use one column: skills and education as sections, one above the other.",
                               right[:2]))
            break
    if extracted is not None and getattr(extracted, "ocr", False):
        out.append(Finding("layout", "problem", "Text had to be recognised from a picture",
                           "The PDF holds images of text; an ATS may read nothing.", "Export a text PDF."))
    return out


# -- the checks ---------------------------------------------------------------------------

def _contact(parts: Parts) -> list[Finding]:
    from ml.tailoring.structure import EMAIL, PHONE
    out = []
    top = parts.contact_text or "\n".join(parts.text.splitlines()[:6])
    if not EMAIL.search(top):
        sev = "problem" if not EMAIL.search(parts.header_footer) else "warning"
        out.append(Finding("contact", sev, "No email address at the top",
                           "A recruiter who wants to call you looks at the top of page one, and so does a parser.",
                           "Add your email under your name."))
    if not any(sum(c.isdigit() for c in p) >= 10 for p in PHONE.findall(top)):
        out.append(Finding("contact", "warning", "No phone number at the top", "",
                           "Add a phone number under your name."))
    if not re.search(r"linkedin\.com|linkedin", top, re.I):
        out.append(Finding("contact", "tip", "No LinkedIn profile",
                           "Most recruiters check it before calling.", "Add linkedin.com/in/your-name."))
    if re.match(r"\s*[A-Z]{2},\s*(USA|US|United States)\b", top.splitlines()[1] if len(top.splitlines()) > 1 else "") \
            or re.search(r"(^|\n|\|)\s*[A-Z]{2},\s*(USA|US)\b", top):
        out.append(Finding("contact", "tip", "Only a state, no city",
                           "Many ATS filter candidates by distance from a city; a state alone can miss the cut.",
                           "Add the city: “Denver, CO”. Keep “Open to Relocate” if it is true."))
    elif not re.search(r"\b[A-Z][a-zA-Z .'-]+,\s*[A-Z]{2}\b|\b[A-Z][a-z]+,\s*[A-Z][a-z]+", top):
        out.append(Finding("contact", "tip", "No city on the resume",
                           "Many ATS filter by location; without one you can miss the cut.",
                           "Add City, State (no street address)."))
    if re.search(r"\b\d{1,5}\s+[A-Z][a-z]+\s+(street|st|avenue|ave|road|rd|lane|ln|drive|dr|blvd)\b", top, re.I):
        out.append(Finding("contact", "warning", "A full street address",
                           "No employer needs it before an offer, and it is personal data on every copy you send.",
                           "Keep only City, State."))
    if not parts.name:
        out.append(Finding("contact", "problem", "No name found at the top", "", "Put your full name on the first line."))
    return out


def _sections(parts: Parts) -> list[Finding]:
    out = []
    kinds = {k for _, k in parts.headings if k}
    if sum(k == "skills" for _, k in parts.headings) > 1:
        out.append(Finding("sections", "tip", "More than one skills section",
                           "The reader meets the same tools twice before the first job.",
                           "Merge them into one Skills section, grouped by label.",
                           [h for h, k in parts.headings if k == "skills"]))
    if "experience" not in kinds and "projects" not in kinds:
        out.append(Finding("sections", "problem", "No Experience section found",
                           "An ATS looks for a heading such as Experience or Work Experience to find your jobs.",
                           "Head the jobs EXPERIENCE (or WORK EXPERIENCE)."))
    if "education" not in kinds:
        out.append(Finding("sections", "warning", "No Education section found",
                           "Many applications are screened on a degree; the ATS looks for this heading.",
                           "Add EDUCATION, even with one line."))
    if "skills" not in kinds:
        out.append(Finding("sections", "problem", "No Skills section found",
                           "A skills section is where an ATS extracts and scores your tools first.",
                           "Add SKILLS with your tools in groups (Cloud: AWS, Azure)."))
    if "summary" not in kinds:
        out.append(Finding("sections", "tip", "No summary",
                           "Two or three lines at the top tell the reader who you are before they read the jobs.",
                           "Add SUMMARY: your title, years, and two strengths with a number."))
    odd = [h for h, k in parts.headings if not k]
    if odd:
        out.append(Finding("sections", "warning", "Headings an ATS may not recognise",
                           "Parsers find sections by their usual names.",
                           "Use the usual names: Summary, Experience, Skills, Education, Projects, Certifications.",
                           odd[:6]))
    named = [h for h, k in parts.headings if k]
    if len(set(h.lower() for h in named)) < len(named):
        dup = sorted({h for h in named if [x.lower() for x in named].count(h.lower()) > 1})
        out.append(Finding("sections", "tip", "The same heading appears twice", "", "Merge them into one.", dup))
    return out


def is_list_line(line: str) -> bool:
    """"Environment: Azure, AWS, ..." or "AWS (EKS, S3) | Azure (AKS)": a list of tools, not a sentence."""
    text = _clean(line)
    return (bool(re.match(r"(environment|tools|tech stack|technologies|stack)\b", text, re.I))
            or text.count(",") + text.count("|") >= 6 and not re.search(r"\b(by|to|for|with)\b.*\d", text))


def bullet_findings(bullets: list[str]) -> list[Finding]:
    """The bullet checks alone: the builder runs these on each job as it is typed."""
    out = []
    lines = [_clean(b) for b in bullets if _clean(b) and not is_list_line(b)]
    if not lines:
        return out
    weak = [l for l in lines if l.lower().startswith(WEAK_STARTS)]
    if weak:
        out.append(Finding("impact", "problem", f"{len(weak)} bullet(s) describe duties, not results",
                           "“Responsible for …” says what the job was. A reader wants what you did and what changed.",
                           "Start with a verb for what you did (Built, Cut, Led), then the result.", weak[:5]))
    pron = [l for l in lines if PRONOUNS.search(l)]
    if pron:
        out.append(Finding("impact", "warning", f"{len(pron)} bullet(s) say “I”, “my” or “we”",
                           "Resumes are written without pronouns; “we” also hides what you did yourself.",
                           "Drop the pronoun and start with the verb.", pron[:5]))
    no_verb = [l for l in lines if l not in weak and l not in pron
               and _first_word(l) not in ACTION_VERBS and not _first_word(l).endswith("ed")]
    if no_verb and len(no_verb) >= max(2, len(lines) // 5):
        out.append(Finding("impact", "tip", f"{len(no_verb)} bullet(s) do not start with an action verb",
                           "A verb first makes a line scan as something you did.",
                           "Lead with the verb: “Migrated …”, “Automated …”.", no_verb[:5]))
    plain = [l for l in lines if not NUMBER.search(l)]
    share = round(100 * (len(lines) - len(plain)) / len(lines))
    if share < 50:
        out.append(Finding("impact", "problem" if share < 25 else "warning",
                           f"Only {share}% of bullets show a number",
                           "A number (users, time saved, cost, size, how many) is what makes a claim believable. "
                           "Aim for at least half.",
                           "Add the real figure where you know it: how many, how much, how fast, how often. "
                           "Never estimate one you cannot back up in an interview.", plain[:6]))
    long = [l for l in lines if len(l.split()) > 40]
    if long:
        out.append(Finding("impact", "warning", f"{len(long)} bullet(s) over 40 words",
                           "Past two lines a bullet stops being read.", "Split it, or cut to the result.", long[:4]))
    short = [l for l in lines if len(l.split()) < 6]
    if short:
        out.append(Finding("impact", "tip", f"{len(short)} bullet(s) under 6 words",
                           "Too short to show what was done and what came of it.",
                           "Add what changed because of it.", short[:4]))
    starts = {}
    for l in lines:
        starts.setdefault(_first_word(l), []).append(l)
    repeated = {w: ls for w, ls in starts.items() if w and len(ls) >= 4}
    for w, ls in sorted(repeated.items(), key=lambda kv: -len(kv[1]))[:3]:
        out.append(Finding("impact", "tip", f"“{w.capitalize()}” starts {len(ls)} bullets",
                           "The same opening verb again and again reads as a template.",
                           "Vary the verbs: Built, Designed, Automated, Migrated, Cut, Led.", ls[:3]))
    cliche = sorted({c for l in lines for c in CLICHES if c in l.lower()})
    if cliche:
        out.append(Finding("impact", "tip", "Phrases that show nothing: " + ", ".join(cliche[:5]),
                           "Every resume says them, so a reader skips them.", "Show it with a result instead."))
    vague = [l for l in lines if re.search(r"\b(" + "|".join(map(re.escape, VAGUE)) + r")\b", l, re.I)]
    if len(vague) >= 3:
        out.append(Finding("impact", "tip", f"{len(vague)} bullet(s) use “various”, “several”, “etc.”",
                           "Vague counts read as filler.", "Name them, or give the number.", vague[:3]))
    return out


def _summary_findings(parts: Parts) -> list[Finding]:
    out = []
    summary = " ".join(parts.sections.get("summary", []))
    if summary:
        words = len(summary.split())
        if words > 120:
            out.append(Finding("impact", "tip", f"The summary is {words} words",
                               "Past about 80 words a summary is skimmed over.", "Keep the two or three strongest lines."))
        if PRONOUNS.search(summary):
            out.append(Finding("polish", "tip", "The summary uses “I” or “my”", "",
                               "Write it without pronouns: “DevOps engineer with 8 years …”."))
        cl = sorted({c for c in CLICHES if c in summary.lower()})
        if cl:
            out.append(Finding("impact", "tip", "The summary leans on stock phrases: " + ", ".join(cl[:4]), "",
                               "Swap each for a fact: a tool, a scale, a result."))
    return out


def _dates(parts: Parts) -> list[Finding]:
    out = []
    jobs = parts.jobs
    undated = [l for l in parts.sections.get("experience", [])[:1]] if not jobs and parts.sections.get("experience") else []
    if undated:
        out.append(Finding("dates", "problem", "No dates found on the jobs",
                           "An ATS works out your years of experience from the dates; with none it may score you as junior.",
                           "Put Month Year – Month Year (or Present) on each job line."))
    text = "\n".join(j["header"] for j in jobs) or parts.text
    styles = sum(bool(p.search(text)) for p in (MONTH_WORD, MONTH_NUM))
    if styles > 1:
        out.append(Finding("dates", "tip", "Two date formats (“Mar 2021” and “03/2021”)", "",
                           "Use one format throughout; “Mar 2021” reads best."))
    today = date.today().strftime("%Y-%m")
    future = [j["header"] for j in jobs if (j["start"] and j["start"] > today) or (j["end"] and j["end"] > today)]
    if future:
        out.append(Finding("dates", "problem", "Dates in the future", "", "Check the years.", future[:3]))
    backwards = [j["header"] for j in jobs if j["start"] and j["end"] and j["start"] > j["end"]]
    if backwards:
        out.append(Finding("dates", "problem", "A job ends before it starts", "", "Check the dates.", backwards[:3]))
    starts = [j["start"] for j in jobs if j["start"]]
    if len(starts) >= 2 and starts != sorted(starts, reverse=True):
        out.append(Finding("dates", "warning", "Jobs are not newest first",
                           "Readers and parsers expect the current job at the top.", "Order the jobs newest first."))
    spans = sorted(((j["start"], j["end"] or (today if j["present"] else None)) for j in jobs
                    if j["start"]), key=lambda s: s[0])
    gaps = []
    reach = None
    for start, end in spans:
        if reach and start > reach and _months(reach, start) > 6:
            gaps.append(f"{_label(reach)} to {_label(start)} ({_months(reach, start)} months)")
        if end and (reach is None or end > reach):
            reach = end
    if gaps:
        out.append(Finding("dates", "tip", "Gaps between jobs",
                           "Not a problem in itself; a reader will ask, so have the one-line answer ready "
                           "(study, family, a move, a contract).", "Optionally note it in a line under the earlier job.",
                           gaps[:4]))
    return out


def _months(a: str, b: str) -> int:
    return (int(b[:4]) - int(a[:4])) * 12 + int(b[5:7]) - int(a[5:7])


def _label(ym: str) -> str:
    import calendar
    return f"{calendar.month_abbr[int(ym[5:7])]} {ym[:4]}"


def years(parts: Parts) -> float | None:
    if parts.worked_years is not None:
        return parts.worked_years
    starts = [j["start"] for j in parts.jobs if j["start"]]
    if not starts:
        return None
    first = min(starts)
    today = date.today()
    return round(today.year - int(first[:4]) + (today.month - int(first[5:7])) / 12, 1)


def _length(parts: Parts) -> list[Finding]:
    from ml.resume.fit import length_for_years
    out = []
    pages, yrs = parts.pages, years(parts)
    if pages and yrs is not None:
        want = length_for_years(yrs)
        top = {"1": 1, "2": 2, "3-4": 4}[want]
        if pages > top:
            out.append(Finding("length", "warning", f"{pages} pages for {yrs:g} years of experience",
                               f"About {want} page(s) is usual for this much experience.",
                               "Cut the oldest jobs to title and dates, and the bullets that repeat."))
        elif want == "3-4" and pages == 1:
            out.append(Finding("length", "tip", f"One page for {yrs:g} years",
                               "A senior career fits more; tools and scale are what get it read.",
                               "Two pages or more is fine at this level."))
    words = len(parts.text.split())
    if words < 200:
        out.append(Finding("length", "warning", f"Only {words} words",
                           "Too little for an ATS to match against most postings.",
                           "Add what you did in each job, with results."))
    return out


def _polish(parts: Parts) -> list[Finding]:
    out = []
    doubled = sorted({m.group(0) for m in re.finditer(r"\b(\w+)\s+\1\b", parts.text, re.I)
                      if not m.group(1).isdigit() and m.group(1).lower() not in ("that", "had")})
    if doubled:
        out.append(Finding("polish", "warning", "A doubled word", "", "Delete the repeat.", doubled[:4]))
    lines = [_clean(b) for b in parts.bullets if _clean(b) and not is_list_line(b)]
    if len(lines) >= 4:
        dots = sum(l.endswith(".") for l in lines)
        if 0 < dots < len(lines) and min(dots, len(lines) - dots) >= 2:
            out.append(Finding("polish", "tip", f"{dots} bullets end with a full stop, {len(lines) - dots} do not", "",
                               "Pick one; most resumes end every bullet with a full stop."))
    if re.search(r"[A-Za-z,]  +[A-Za-z]", parts.text):
        out.append(Finding("polish", "tip", "Double spaces between words", "", "Replace them with one space."))
    if re.search(r"\b(utilize|utilized|utilizing)\b", parts.text, re.I):
        out.append(Finding("polish", "tip", "“Utilized”", "It is “used” in more letters.", "Write “used”."))
    return out


def check(parts: Parts) -> Report:
    findings = (_contact(parts) + _sections(parts) + bullet_findings(parts.bullets) + _summary_findings(parts)
                + _dates(parts) + _length(parts) + parts.layout + _polish(parts))
    order = {"problem": 0, "warning": 1, "tip": 2}
    findings.sort(key=lambda f: (order[f.severity], list(WEIGHTS).index(f.category)))
    categories = []
    total = 0.0
    for cat, share in WEIGHTS.items():
        lost = sum(COST[f.severity] for f in findings if f.category == cat)
        # Three problems take most of a category; never below a tenth of it.
        got = share * max(0.1, 1 - lost) if lost else share
        total += got
        mine = [f for f in findings if f.category == cat]
        categories.append({"key": cat, "label": LABELS[cat], "score": round(got), "out_of": share,
                           "problems": sum(f.severity == "problem" for f in mine),
                           "warnings": sum(f.severity == "warning" for f in mine),
                           "tips": sum(f.severity == "tip" for f in mine)})
    lines = [_clean(b) for b in parts.bullets if _clean(b) and not is_list_line(b)]
    stats = {"bullets": len(lines),
             "with_numbers": sum(bool(NUMBER.search(l)) for l in lines),
             "action_verbs": sum(_first_word(l) in ACTION_VERBS or _first_word(l).endswith("ed") for l in lines),
             "words": len(parts.text.split()), "pages": parts.pages, "years": years(parts),
             "jobs": len(parts.jobs), "sections": [h for h, _ in parts.headings]}
    return Report(round(total), categories, findings, stats)


def check_file(path: Path | str) -> Report:
    return check(from_file(path))


# -- rewriting a weak bullet, on request ------------------------------------------------------

REWRITE_SYSTEM = """You rewrite resume bullets so they read as results, not duties.
For each bullet you get, write one better version. Reply with JSON only:
{"bullets": [{"original": "the bullet as given", "rewrite": "the better version",
              "ask": "a question for the person when a number would make it stronger, or empty"}]}
Rules:
- Use ONLY what the bullet and the context say. Never add a tool, number, employer, team size or result
  that is not written there. If the bullet has no number, do not invent one: write it without, and put
  the question that would get the real number in "ask" ("How many services did this cover?").
- Start with a strong past-tense action verb (present tense only if the context says it is the current job).
- No "I", "my", "we". No stock phrases. One sentence, under 30 words. Plain words, no long dashes."""


def rewrite_bullets(bullets: list[str], context: str, cfg, ask=None) -> list[dict]:
    """Better versions of weak bullets, each checked: a number or tool not in the source is flagged."""
    from backend.apply.essays import numbers_not_on
    if not bullets:
        return []
    if ask is None:
        from ml.tailoring.review import _ask as ask
    listed = "\n".join(f"- {_clean(b)}" for b in bullets[:12])
    body = ask(REWRITE_SYSTEM, f"CONTEXT (the rest of the resume):\n{context[:12000]}\n\nBULLETS:\n{listed}", cfg)
    out = []
    source = context + "\n" + listed
    for item in body.get("bullets") or []:
        if not isinstance(item, dict) or not str(item.get("rewrite") or "").strip():
            continue
        rewrite = " ".join(str(item["rewrite"]).split())
        out.append({"original": str(item.get("original") or "").strip(), "rewrite": rewrite,
                    "ask": str(item.get("ask") or "").strip(),
                    "unsupported_numbers": numbers_not_on(rewrite, source),
                    "unsupported_tools": tools_not_on(rewrite, source)})
    return out


def tools_not_on(text: str, source: str) -> list[str]:
    """Named tools in `text` that `source` never mentions: a claim the person did not make."""
    from ml.resume.scorer import SKILL_TERMS, tokenize
    have = {t.strip(".-") for t in tokenize(source or "")}
    generic = {"api", "apis", "backend", "frontend", "fullstack", "architecture", "algorithms", "debugging",
               "profiling", "mentoring", "leadership", "scalability", "latency", "throughput", "caching",
               "observability", "ci", "cd", "devops", "sre", "etl", "rest"}
    words = {t.strip(".-") for t in tokenize(text or "")}
    return sorted(t for t in words if t in SKILL_TERMS and t not in generic and t not in have)
