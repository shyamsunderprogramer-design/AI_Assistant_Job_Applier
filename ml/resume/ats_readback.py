"""Read a resume back the way an ATS does, and say what it would get wrong.

An applicant tracking system does not see the page. It sees the text the file gives it, in the order
the file gives it, and parsers differ -- so each file is read by more than one, the way different
systems read it:

  PDF    pypdf, PDFium (Chrome's PDF engine) and pdfminer (the engine under many resume parsers);
  .docx  python-docx, and the raw XML of the document read paragraph by paragraph, which is all the
         simplest parsers do.

Each reading is checked against what was meant to be on the page (the planned blocks):

  * the name first, then an email and a phone in the body of the page;
  * every section heading on a line of its own, and in the planned order;
  * every job's title line with its dates beside it or on the next line;
  * every bullet, skills line and education line whole and in one piece, in order -- a page read
    across two columns at once breaks lines apart;
  * nothing a parser chokes on: ligatures ("ﬁ"), unmapped glyphs, and in a .docx no tables, text
    boxes, pictures or newspaper columns, and nothing in the page header or footer, which many
    parsers skip.

A template is offered for applying only when every reader gets every line right (verify_template).
"""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
_BAD_CHARS = re.compile(r"[ﬀ-ﬆ-�]|\(cid:\d+\)")
_DATES = re.compile(r"\b(?:19|20)\d{2}\b|\bpresent\b|\bcurrent\b", re.I)


@dataclass
class Finding:
    reader: str
    kind: str
    detail: str


@dataclass
class Reading:
    reader: str
    text: str
    findings: list[Finding] = field(default_factory=list)


# -- the readers -----------------------------------------------------------------------------------

def _pypdf(pdf: Path) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)


def _pdfium(pdf: Path) -> str:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf))
    try:
        return "\n".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))
    finally:
        doc.close()


def _pdfminer(pdf: Path) -> str:
    from pdfminer.high_level import extract_text
    return extract_text(str(pdf))


def _python_docx(docx: Path) -> str:
    from docx import Document
    return "\n".join(p.text for p in Document(str(docx)).paragraphs)


def _xml_paragraphs(xml: str) -> list[str]:
    out = []
    for para in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        bits = re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>|<w:(tab|br)\b[^>]*/>", para, re.S)
        text = "".join(t if t else ("\t" if kind == "tab" else "\n") for t, kind in bits)
        out.append(unicodedata.normalize("NFC", re.sub(r"&amp;", "&", re.sub(r"&lt;", "<", re.sub(r"&gt;", ">", text)))))
    return out


def _raw_docx(docx: Path) -> str:
    with zipfile.ZipFile(docx) as z:
        return "\n".join(_xml_paragraphs(z.read("word/document.xml").decode("utf-8", "replace")))


PDF_READERS = {"pypdf": _pypdf, "PDFium": _pdfium, "pdfminer": _pdfminer}
DOCX_READERS = {"python-docx": _python_docx, "raw XML": _raw_docx}


# -- comparing a reading with the plan ---------------------------------------------------------------

def _compact(text: str) -> str:
    """Letters and digits only, lower case: line wraps, spacing and punctuation readers differ on gone."""
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKC", text).lower())


def _lines(text: str) -> list[str]:
    return [l.strip() for l in text.splitlines() if l.strip()]


def _check(reading: Reading, blocks: list[tuple[str, str]]) -> None:
    from ml.resume.writer import _split_dates
    from ml.resume.parser import strip_bullet
    text, who = reading.text, reading.reader
    flat = _compact(text)
    lines = _lines(text)
    # Where each line ends in `flat` (which is the lines' compact text run together).
    ends, total = [], 0
    for l in lines:
        total += len(_compact(l))
        ends.append(total)
    say = lambda kind, detail: reading.findings.append(Finding(who, kind, detail))

    if _BAD_CHARS.search(text):
        sample = _BAD_CHARS.search(text).group(0)
        say("unreadable characters", f"the text holds {sample!r}, which parsers read as garbage")

    name = next((t for k, t in blocks if k == "name"), "")
    if name and (not lines or _compact(name) not in _compact(lines[0])):
        say("name", f"the first line read is “{(lines[0] if lines else '')[:50]}”, not the name")

    contact = " ".join(t for k, t in blocks if k == "contact")
    email = _EMAIL.search(contact)
    if email and email.group(0).lower() not in text.lower():
        say("contact", "the email address is not in the text")
    digits = re.sub(r"\D", "", contact)
    phone = re.search(r"\d{10,}", digits)
    if phone and phone.group(0)[-7:] not in re.sub(r"\D", "", text):
        say("contact", "the phone number is not in the text")

    at = 0
    for kind, planned in blocks:
        if kind == "heading":
            want = _compact(planned)
            own = [i for i, l in enumerate(lines) if _compact(l) == want]
            if not own:
                ran = next((l for l in lines if _compact(l).startswith(want) or _compact(l).endswith(want)), "")
                say("heading", f"“{planned}” is not on a line of its own"
                    + (f" (read as “{ran[:60]}”)" if ran else ""))
            spot = flat.find(want, at)
            if spot < 0:
                say("order", f"“{planned}” is out of order or missing")
            else:
                at = spot
        elif kind in ("bullet", "text", "skill", "degree", "school", "role"):
            body = strip_bullet(planned)
            dates = ""
            if kind in ("role", "degree"):
                body, dates = _split_dates(body)
            want = _compact(body)
            if not want:
                continue
            spot = flat.find(want, at)
            if spot < 0:
                spot = flat.find(want)
                what = {"role": "job title line", "degree": "degree line", "school": "school line",
                        "skill": "skills line"}.get(kind, "line")
                say("broken line" if spot < 0 else "order",
                    f"the {what} “{body[:60]}” is " + ("not read in one piece" if spot < 0 else "read out of order"))
                if spot < 0:
                    continue
            if dates:
                # On the title's own line or the next one: dates read anywhere else belong to nobody.
                last = next((i for i, e in enumerate(ends) if e >= spot + len(want)), len(ends) - 1)
                upto = ends[min(last + 1, len(ends) - 1)] if ends else len(flat)
                if _compact(dates) not in flat[spot:upto]:
                    say("dates", f"the dates “{dates}” are not beside “{body[:40]}”")
            at = spot + len(want)


def _docx_structure(docx: Path) -> list[Finding]:
    """What many parsers lose in a .docx: tables, text boxes, pictures, columns, a header or footer."""
    out = []
    with zipfile.ZipFile(docx) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
        for tag, what in (("<w:tbl>", "a table"), ("<w:txbxContent", "a text box"), ("<w:drawing", "a picture"),
                          ("<w:pict", "a picture")):
            if tag in xml:
                out.append(Finding("structure", "layout", f"the document has {what}, which many parsers skip"))
        if re.search(r'<w:cols\b[^>]*w:num="([2-9])"', xml):
            out.append(Finding("structure", "layout", "the page is set in columns, read across by many parsers"))
        for name in z.namelist():
            if re.match(r"word/(header|footer)\d*\.xml$", name):
                words = "".join(_xml_paragraphs(z.read(name).decode("utf-8", "replace"))).strip()
                if words:
                    out.append(Finding("structure", "layout",
                                       f"the page {name.split('/')[1].rstrip('0123456789.xml')} holds “{words[:40]}”, "
                                       "which many parsers skip"))
    return out


def read_back(pdf: Path | None, docx: Path | None, blocks: list[tuple[str, str]]) -> list[Reading]:
    """Every reader's reading of the files, each with what it got wrong."""
    readings = []
    for files, readers in ((pdf, PDF_READERS), (docx, DOCX_READERS)):
        if files is None or not Path(files).exists():
            continue
        for name, read in readers.items():
            try:
                reading = Reading(name, read(Path(files)))
            except Exception as exc:                       # a reader that cannot open the file is a finding
                reading = Reading(name, "")
                reading.findings.append(Finding(name, "unreadable", f"could not read the file ({type(exc).__name__})"))
                readings.append(reading)
                continue
            _check(reading, blocks)
            readings.append(reading)
    if docx is not None and Path(docx).exists():
        structure = Reading("structure", "")
        structure.findings = _docx_structure(Path(docx))
        readings.append(structure)
    return readings


def checks(pdf: Path | None, docx: Path | None, blocks: list[tuple[str, str]]):
    """The findings as the tailoring page's checks (ml/tailoring/documents.Check)."""
    from ml.tailoring.documents import Check
    out = []
    for reading in read_back(pdf, docx, blocks):
        file = "docx" if reading.reader in DOCX_READERS or reading.reader == "structure" else "pdf"
        for f in reading.findings[:6]:
            out.append(Check(file, None, "ATS read", f"{f.reader}: {f.detail}", "problem"))
    return out


# -- a template, proven or not --------------------------------------------------------------------------

def verify_template(template: str, base, header: list[str] | None, folder: Path) -> dict:
    """Draw the resume in one template and read it back with every reader: {template, ok, readers}."""
    from ml.resume.pdf import write_pdf
    from ml.resume.tailor import TailorResult
    from ml.resume.writer import plan_document, write_tailored_resume
    folder.mkdir(parents=True, exist_ok=True)
    result = TailorResult(summary=None, template=template)
    docx, pdf = folder / f"{template}.docx", folder / f"{template}.pdf"
    write_tailored_resume(base, result, docx, header_lines=header)
    write_pdf(base, result, pdf, header_lines=header)
    blocks = plan_document(base, result, header)
    readings = read_back(pdf, docx, blocks)
    return {"template": template, "ok": not any(r.findings for r in readings),
            "readers": {r.reader: [f"{f.kind}: {f.detail}" for f in r.findings] for r in readings}}
