#!/usr/bin/env python3
"""Build a tailored resume .docx in the skill's format (references/format.md).

Usage:
    python build_resume.py content.json out.docx [--pdf]

Needs python-docx (`pip install python-docx`). --pdf also writes a PDF if
LibreOffice (`soffice`) or Microsoft Word (via docx2pdf) is available.

content.json:
{
  "name": "Jordan Example",
  "contact": ["Austin, TX (Open to Relocate)", "+1 555 010 0199", "jordan@example.com",
              {"text": "LinkedIn", "url": "https://linkedin.com/in/jordan-example"}],
  "title": "Senior Platform Engineer · Azure & AWS · Terraform · Kubernetes",
  "keywords": ["kubernetes", "terraform", "ci/cd", "infrastructure as code"],
  "sections": [
    {"heading": "Summary", "paragraphs": ["..."]},
    {"heading": "Highlights", "bullets": ["..."]},
    {"heading": "Core Competencies", "rows": [["Cloud", "AWS; Azure; GCP."]]},
    {"heading": "Experience", "jobs": [
        {"title": "Senior DevOps Engineer, Acme, USA", "dates": "Feb 2025 – Present",
         "groups": [{"name": "Cloud Infrastructure", "bullets": ["..."]}],
         "bullets": ["bullets with no group go here"]}]},
    {"heading": "Projects", "projects": [{"name": "...", "description": "..."}]},
    {"heading": "Certifications", "certs": [{"name": "AWS Solutions Architect", "rest": "Amazon, 2023"}]},
    {"heading": "Education", "schools": [{"school": "Example University", "degree": "Master of Science",
                                          "dates": "2023 – 2024"}]}
  ]
}

Sections are put in the format's order whatever order they arrive in, and
empty ones are left out. Keywords are bolded under the skill's rules: the
first use of each on a line, whole product names, at most three a line.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Cambria"
BODY = Pt(11)
BLACK = RGBColor(0, 0, 0)
MARGIN_IN = 0.65
TEXT_WIDTH = Inches(8.5 - 2 * MARGIN_IN)
BULLET_AT, TEXT_AT = Inches(0.15), Inches(0.30)
PER_LINE = 3

ORDER = (("summary", "objective", "profile"), ("highlight", "impact", "achievement"),
         ("expertise", "practice"), ("skill", "competenc", "technolog"),
         ("experience", "employment", "work history"), ("project",),
         ("certif", "license"), ("education",))


# -- keyword emphasis ------------------------------------------------------------

def spans(text: str, keywords: list[str], limit: int = PER_LINE) -> list[tuple[str, bool]]:
    """(piece, bold) runs: each keyword's first use on the line, whole names, at most `limit`."""
    if not keywords or not text:
        return [(text, False)]
    found: list[tuple[int, int]] = []
    for term in sorted(keywords, key=len, reverse=True):
        m = re.search(rf"(?<![\w/]){re.escape(term)}(?![\w/])", text, re.I)
        if not m:
            continue
        start, end = m.start(), m.end()
        while nxt := re.match(r" [A-Z][\w+#-]*(?![\w/])", text[end:]):   # "Azure Key Vault"
            end += len(nxt.group(0))
        for a, b in [x for x in found if x[0] < end and start < x[1]]:   # "Azure DevOps"
            found.remove((a, b))
            start, end = min(a, start), max(b, end)
        found.append((start, end))
    found = sorted(found)[:limit]
    out, at = [], 0
    for a, b in found:
        if a > at:
            out.append((text[at:a], False))
        out.append((text[a:b], True))
        at = b
    if at < len(text):
        out.append((text[at:], False))
    return out


# -- paragraphs -----------------------------------------------------------------

def _run(para, text: str, bold: bool = False, size=BODY):
    run = para.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = size
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    return run


def _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), 1.0
    f.keep_with_next = keep_next
    return p


def _emphasised(p, text: str, keywords: list[str], lead: str = "") -> None:
    for i, (piece, bold) in enumerate(spans(text, keywords)):
        _run(p, (lead if i == 0 else "") + piece, bold)


def _heading(doc, text: str) -> None:
    p = doc.add_heading(level=1)
    f = p.paragraph_format
    f.space_before, f.space_after, f.keep_with_next = Pt(8), Pt(4), True
    _run(p, text.upper(), bold=True, size=Pt(12))
    # A hairline under the heading: a paragraph border, invisible to parsers.
    borders, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "2"), ("w:color", "BFBFBF")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p._p.get_or_add_pPr().append(borders)


def _dated(doc, left: str, dates: str, bold: bool) -> None:
    """Title left, dates on a right tab at the margin (two spaces keep them apart in text)."""
    p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _run(p, left + ("  \t" if dates else ""), bold)
    if dates:
        _run(p, dates, bold)


def _bullet(doc, text: str, keywords: list[str]) -> None:
    p = _para(doc)
    f = p.paragraph_format
    f.left_indent, f.first_line_indent = TEXT_AT, BULLET_AT - TEXT_AT
    f.tab_stops.add_tab_stop(TEXT_AT)
    f.keep_together = True
    _emphasised(p, text.lstrip("•-* ").strip(), keywords, lead="•\t")


def _link(p, text: str, url: str) -> None:
    rel = p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT)
    rpr.append(fonts)
    for tag, val in (("w:color", "000000"), ("w:u", "single"), ("w:sz", "22")):
        el = OxmlElement(tag); el.set(qn("w:val"), val); rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; r.append(t)
    link.append(r)
    p._p.append(link)


# -- sections -------------------------------------------------------------------

def _rank(heading: str) -> int:
    low = heading.lower()
    return next((i for i, words in enumerate(ORDER) if any(w in low for w in words)), len(ORDER) - 1)


def build(content: dict, out: Path) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = FONT, BODY, BLACK
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(MARGIN_IN)
    keys = [k.lower() for k in content.get("keywords") or []]

    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=3)
    _run(p, content["name"], bold=True, size=Pt(18))
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=3)
    for i, item in enumerate(content.get("contact") or []):
        if i:
            _run(p, " · ")
        if isinstance(item, dict) and item.get("url"):
            _link(p, item.get("text") or item["url"], item["url"])
        else:
            _run(p, item["text"] if isinstance(item, dict) else str(item))
    if content.get("title"):
        p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=6)
        _run(p, content["title"], bold=True)

    sections = sorted(content.get("sections") or [], key=lambda s: _rank(s.get("heading", "")))
    for sec in sections:
        body = [k for k in ("paragraphs", "bullets", "rows", "jobs", "projects", "certs", "schools")
                if sec.get(k)]
        if not body:
            continue                                   # empty sections are left out
        _heading(doc, sec["heading"])
        for text in sec.get("paragraphs") or []:
            _emphasised(_para(doc), text, keys)
        for text in sec.get("bullets") or []:
            _bullet(doc, text, keys)
        for label, items in sec.get("rows") or []:
            p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT)
            _run(p, f"{label}: ", bold=True)
            _emphasised(p, items, keys)
        for job in sec.get("jobs") or []:
            _dated(doc, job["title"], job.get("dates", ""), bold=True)
            for text in job.get("bullets") or []:
                _bullet(doc, text, keys)
            for group in job.get("groups") or []:
                _run(_para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=2, keep_next=True), group["name"], bold=True)
                for text in group.get("bullets") or []:
                    _bullet(doc, text, keys)
        for proj in sec.get("projects") or []:
            _run(_para(doc, WD_ALIGN_PARAGRAPH.LEFT, keep_next=True), proj["name"], bold=True)
            if proj.get("description"):
                _emphasised(_para(doc), proj["description"], keys)
        for cert in sec.get("certs") or []:
            p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT)
            _run(p, cert["name"], bold=True)
            if cert.get("rest"):
                _run(p, f", {cert['rest']}")
        for school in sec.get("schools") or []:
            _dated(doc, school.get("degree", ""), school.get("dates", ""), bold=False)
            _run(_para(doc, WD_ALIGN_PARAGRAPH.LEFT), school["school"], bold=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def to_pdf(docx: Path) -> Path | None:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if Path(soffice).exists():
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=False, capture_output=True, timeout=120)
    else:
        try:
            from docx2pdf import convert                       # uses Microsoft Word
            convert(str(docx), str(docx.with_suffix(".pdf")))
        except Exception:
            return None
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    out = build(json.loads(Path(sys.argv[1]).read_text()), Path(sys.argv[2]))
    print(f"Wrote {out}")
    if "--pdf" in sys.argv:
        pdf = to_pdf(out)
        print(f"Wrote {pdf}" if pdf else "No PDF: install LibreOffice, or Word with `pip install docx2pdf`.")
