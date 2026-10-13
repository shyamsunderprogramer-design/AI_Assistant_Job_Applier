"""The resume health check: written and built well, with no job posting."""

from docx import Document

from ml.resume import health


def _parts(**kw):
    base = dict(name="Jane Doe", contact_text="Jane Doe\nDenver, CO · 555-010-2030 · jane@example.com · linkedin.com/in/jane",
                headings=[("Summary", "summary"), ("Skills", "skills"), ("Experience", "experience"),
                          ("Education", "education")],
                bullets=["Cut build time from 25 to 8 minutes by moving 40 repositories to GitHub Actions.",
                         "Led a team of 4 engineers through the migration of 12 services to Kubernetes."],
                jobs=[{"header": "Engineer, Acme", "start": "2020-01", "end": None, "present": True, "dates": ""}],
                text="word " * 400, pages=1)
    base.update(kw)
    return health.Parts(**base)


def test_a_sound_resume_scores_high():
    report = health.check(_parts())
    assert report.score >= 95, [f.title for f in report.findings]


def test_duty_bullets_pronouns_and_no_numbers_are_found_with_their_lines():
    found = health.bullet_findings(["Responsible for the CI pipelines.", "I managed our AWS accounts for the team.",
                                    "Worked on monitoring dashboards in Grafana."])
    titles = " | ".join(f.title for f in found)
    assert "describe duties" in titles and "“I”" in titles and "show a number" in titles
    duties = next(f for f in found if "duties" in f.title)
    assert duties.lines == ["Responsible for the CI pipelines.", "Worked on monitoring dashboards in Grafana."]


def test_a_list_of_tools_is_not_judged_as_a_bullet():
    assert health.is_list_line("Environment: Azure, AWS, Terraform, Jenkins, Docker, Helm, Kubernetes")
    assert not health.bullet_findings(["Environment: Azure, AWS, Terraform, Jenkins, Docker, Helm"])


def test_missing_sections_and_unknown_headings():
    report = health.check(_parts(headings=[("My Career Journey", None), ("Experience", "experience")]))
    titles = [f.title for f in report.findings]
    assert "No Skills section found" in titles and "No Education section found" in titles
    odd = next(f for f in report.findings if f.title.startswith("Headings an ATS"))
    assert odd.lines == ["My Career Journey"]


def test_dates_out_of_order_backwards_and_gaps():
    jobs = [{"header": "A", "start": "2015-01", "end": "2016-01", "present": False, "dates": ""},
            {"header": "B", "start": "2019-01", "end": "2018-01", "present": False, "dates": ""}]
    titles = [f.title for f in health.check(_parts(jobs=jobs)).findings]
    assert "Jobs are not newest first" in titles and "A job ends before it starts" in titles
    assert "Gaps between jobs" in titles


def test_no_contact_details_is_a_problem():
    report = health.check(_parts(contact_text="Jane Doe"))
    assert any(f.title == "No email address at the top" and f.severity == "problem" for f in report.findings)


def test_a_docx_with_a_table_and_contact_in_the_header_is_flagged(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "jane@example.com · 555 010 2030"
    doc.add_paragraph("EXPERIENCE")
    doc.add_table(rows=2, cols=2)
    path = tmp_path / "cv.docx"
    doc.save(path)
    titles = [f.title for f in health.layout_findings(path)]
    assert any("table" in t for t in titles)
    assert "Contact details are in the page header or footer" in titles


def test_a_rewrite_that_adds_a_number_or_tool_is_flagged():
    reply = {"bullets": [{"original": "Worked on deployments.", "rewrite": "Automated 50 deployments with Jenkins.",
                          "ask": "How many deployments?"}]}
    out = health.rewrite_bullets(["Worked on deployments."], "Deployed services on AWS.", None,
                                 ask=lambda s, u, c: reply)
    assert out[0]["unsupported_numbers"] == ["50"] and out[0]["unsupported_tools"] == ["jenkins"]
