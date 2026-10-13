"""Templates: the same words, a different look; every one ATS-safe."""

from docx import Document

from ml.resume.parser import parse_resume
from ml.resume.style import TEMPLATES, style_for
from ml.resume.tailor import TailorResult
from ml.resume.writer import write_tailored_resume


def _resume(tmp_path):
    src = tmp_path / "in.docx"
    d = Document()
    for line in ("Jordan Example", "Austin, TX | jordan@example.com", "Experience",
                 "Engineer, Acme   2019 – 2023", "• Built deployment pipelines for forty services."):
        d.add_paragraph(line)
    d.save(src)
    return parse_resume(src)


def test_each_template_sets_its_own_font_and_heading_colour(tmp_path):
    base = _resume(tmp_path)
    for key, st in TEMPLATES.items():
        out = write_tailored_resume(base, TailorResult(summary=None, template=key), tmp_path / f"{key}.docx")
        doc = Document(str(out))
        assert doc.styles["Normal"].font.name == st.font
        assert str(doc.styles["Heading 1"].font.color.rgb) == st.heading_color
        assert not doc.tables                                   # ATS-safe whatever the look
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Built deployment pipelines" in text


def test_a_render_leaves_the_next_one_untouched(tmp_path):
    base = _resume(tmp_path)
    write_tailored_resume(base, TailorResult(summary=None, template="compact"), tmp_path / "a.docx")
    doc = Document(str(write_tailored_resume(base, TailorResult(summary=None), tmp_path / "b.docx")))
    assert doc.styles["Normal"].font.size.pt == TEMPLATES["classic"].body


def test_an_unknown_template_falls_back_to_the_default():
    assert style_for(TailorResult(summary=None, template="neon")).key == "classic"
    assert style_for(None).key == "classic"


def test_the_pdf_draws_each_template(tmp_path):
    from pypdf import PdfReader
    from ml.resume.pdf import write_pdf
    base = _resume(tmp_path)
    for key in TEMPLATES:
        pdf = write_pdf(base, TailorResult(summary=None, template=key), tmp_path / f"{key}.pdf")
        assert "Built deployment pipelines" in PdfReader(str(pdf)).pages[0].extract_text()


def test_a_title_is_never_left_at_the_foot_of_a_page():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import KeepTogether, Paragraph
    from ml.resume.pdf import _kept_with_next
    title = Paragraph("Leadership", ParagraphStyle("t", keepWithNext=1))
    bullet = KeepTogether([Paragraph("Led teams.", ParagraphStyle("b"))])
    out = _kept_with_next([title, bullet, Paragraph("After.", ParagraphStyle("b"))])
    assert len(out) == 2 and isinstance(out[0], KeepTogether)
    assert out[0]._content[0] is title                     # the title travels with its bullet


def test_sidebar_moves_the_content_right_of_the_headings_and_keeps_one_column(tmp_path):
    from docx.shared import Inches
    base = _resume(tmp_path)
    doc = Document(str(write_tailored_resume(base, TailorResult(summary=None, template="sidebar"),
                                             tmp_path / "s.docx")))
    by_text = {p.text.split("\t")[0].strip(): p for p in doc.paragraphs}
    heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    assert heading.paragraph_format.left_indent in (None, 0)          # the rail
    role = next(p for p in doc.paragraphs if p.text.startswith("Engineer, Acme"))
    assert role.paragraph_format.left_indent == Inches(1.3)            # the work, beside it
    bullet = next(p for p in doc.paragraphs if "Built deployment" in p.text)
    assert bullet.paragraph_format.left_indent == Inches(1.3 + 0.30)
    body = doc.element.body
    from docx.oxml.ns import qn
    assert not doc.tables and not body.findall(".//" + qn("w:txbxContent"))
    cols = [c for c in body.iter(qn("w:cols")) if int(c.get(qn("w:num"), "1") or 1) > 1]
    assert not cols and by_text


def test_minimal_has_no_rules_and_a_left_header(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    base = _resume(tmp_path)
    doc = Document(str(write_tailored_resume(base, TailorResult(summary=None, template="minimal"),
                                             tmp_path / "m.docx")))
    assert not doc.element.body.findall(".//" + qn("w:pBdr"))
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_a_look_changes_the_font_and_the_heading_colour(tmp_path):
    from pypdf import PdfReader
    from ml.resume.pdf import write_pdf
    base = _resume(tmp_path)
    result = TailorResult(summary=None, template="classic", look={"font": "Georgia", "accent": "7A1F2B"})
    doc = Document(str(write_tailored_resume(base, result, tmp_path / "g.docx")))
    assert doc.styles["Normal"].font.name == "Georgia"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "7A1F2B"
    pdf = write_pdf(base, result, tmp_path / "g.pdf")
    fonts = {str(f["/BaseFont"]) for f in PdfReader(str(pdf)).pages[0]["/Resources"]["/Font"].values()
             for f in [f.get_object()]}
    assert any("Georgia" in f for f in fonts), fonts


def test_an_unknown_font_or_colour_is_ignored():
    from ml.resume.style import clean_look
    assert clean_look({"font": "Comic Sans", "accent": "#ff00ff", "x": 1}) == {}
    assert clean_look({"font": "Arial", "accent": "#1f3864"}) == {"font": "Arial", "accent": "1F3864"}
