"""Reading resumes for Resume Tailoring: DOCX, PDF, scanned PDF, and every failure."""

import pytest
from docx import Document
from docx.shared import Pt

from ml.tailoring.reader import ReadError, read_document
from ml.tailoring.structure import build


def make_docx(path, unusual=False):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Jordan Example — Resume"
    d.sections[0].footer.paragraphs[0].text = "Page 1"
    d.add_paragraph("Jordan Example")
    contact = d.add_paragraph("jordan@example.com | +1 555 010 0199 | ")
    from ml.resume.writer import _hyperlink
    _hyperlink(contact, "LinkedIn", "https://linkedin.com/in/jordan-example")
    if unusual:
        h = d.add_paragraph(); r = h.add_run("Professional Experience"); r.bold = True; r.font.size = Pt(14)
    else:
        d.add_heading("Experience", level=1)
    d.add_paragraph("Platform Engineer, Acme Corp, Austin TX     Jan 2021 – Present")
    d.add_paragraph("Built Terraform modules for 200+ AWS accounts.", style="List Bullet")
    d.add_paragraph("Ran Kubernetes 1.27 clusters on EKS.", style="List Bullet")
    d.add_paragraph("Systems Engineer, Globex, Remote     Mar 2017 – Dec 2020")
    d.add_paragraph("Migrated 40 services to Docker.", style="List Number")
    d.add_heading("Skills", level=1)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Cloud", "AWS, Azure"
    table.cell(1, 0).text, table.cell(1, 1).text = "IaC", "Terraform, Ansible"
    d.add_heading("Education", level=1)
    d.add_paragraph("B.S. Computer Science, State University     2013 – 2017")
    d.save(path)
    return path


def test_docx_reads_in_order_with_lists_tables_links_headers(tmp_path):
    x = read_document(make_docx(tmp_path / "r.docx"))
    sources = [b.source for b in x.blocks]
    assert sources[0] == "header" and sources[-1] == "footer"
    assert any(s.startswith("table 1, row 2, cell 2") for s in sources)
    texts = [b.text for b in x.blocks]
    assert texts.index("AWS, Azure") > texts.index("Migrated 40 services to Docker.")   # document order
    assert next(b for b in x.blocks if b.text.startswith("Migrated")).numbered
    assert next(b for b in x.blocks if b.text.startswith("Built")).list_item
    assert any("linkedin.com/in/jordan-example" in l for b in x.blocks for l in b.links)
    assert next(b for b in x.blocks if b.text == "Experience").heading


def test_each_bullet_stays_with_its_own_employer(tmp_path):
    m = build(read_document(make_docx(tmp_path / "r.docx")))
    acme, globex = [e for e in m.entries if e.section == "experience"]
    assert (acme.title, acme.employer) == ("Platform Engineer", "Acme Corp, Austin TX")
    assert acme.start == "2021-01" and acme.present
    assert [m.fact(f).text for f in acme.facts][1:] == ["Built Terraform modules for 200+ AWS accounts.",
                                                         "Ran Kubernetes 1.27 clusters on EKS."]
    assert [m.fact(f).text for f in globex.facts][1:] == ["Migrated 40 services to Docker."]
    assert all(m.fact(f).source.startswith("paragraph") for f in acme.facts)
    assert "Terraform, Ansible" in [m.fact(f).text for f in m.sections["skills"]]


def test_unusual_formatting_still_finds_the_sections(tmp_path):
    m = build(read_document(make_docx(tmp_path / "r.docx", unusual=True)))
    assert len([e for e in m.entries if e.section == "experience"]) == 2


def text_pdf(path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=LETTER)
    lines = ["JORDAN EXAMPLE", "jordan@example.com", "EXPERIENCE",
             "Platform Engineer, Acme Corp     Jan 2021 - Present",
             "- Built Terraform modules for 200 AWS accounts.", "EDUCATION",
             "B.S. Computer Science, State University 2013 - 2017"]
    for n, line in enumerate(lines):
        c.setFont("Helvetica", 14)
        c.drawString(72, 700 - n * 26, line)
    c.save()
    return path


def test_a_text_pdf_reads_with_page_and_line_sources(tmp_path):
    x = read_document(text_pdf(tmp_path / "r.pdf"))
    assert not x.ocr and x.blocks[0].source == "page 1, line 1"
    m = build(x)
    assert [e.employer for e in m.entries if e.section == "experience"] == ["Acme Corp"]


def test_a_scanned_pdf_is_recognised_and_marked_low_confidence(tmp_path):
    pdfium = pytest.importorskip("pypdfium2")
    pytest.importorskip("ocrmac")
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    image = pdfium.PdfDocument(str(text_pdf(tmp_path / "t.pdf")))[0].render(scale=2).to_pil()
    scanned = tmp_path / "scanned.pdf"
    c = canvas.Canvas(str(scanned), pagesize=LETTER)
    c.drawImage(ImageReader(image), 0, 0, *LETTER)
    c.save()
    x = read_document(scanned)
    assert x.ocr and all(b.confidence == "low" for b in x.blocks)
    assert any("recognised" in w for w in x.warnings)
    assert "Terraform" in x.text and "EXPERIENCE" in x.text.upper()


@pytest.mark.parametrize("name,content,kind", [
    ("r.docx", b"this is not a zip file", "damaged"),
    ("r.pdf", b"%PDF-1.4 garbage", "damaged"),
    ("r.png", b"\x89PNG", "unsupported"),
    ("r.docx", b"", "empty"),
])
def test_bad_files_say_what_is_wrong(tmp_path, name, content, kind):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ReadError) as err:
        read_document(path)
    assert err.value.kind == kind and name in str(err.value)


def test_a_missing_file(tmp_path):
    with pytest.raises(ReadError) as err:
        read_document(tmp_path / "gone.docx")
    assert err.value.kind == "missing"


def test_unreadable_dates_and_missing_headings_are_reported(tmp_path):
    from ml.tailoring.reader import Block, Extracted
    blocks = [Block("Jordan Example", "line 1"), Block("Did various things for a while.", "line 2")]
    m = build(Extracted("text", blocks))
    assert any("No section headings" in w for w in m.warnings)
    blocks = [Block("EXPERIENCE", "line 1"), Block("Engineer, Acme   Dec 2022 – Jan 2020", "line 2"),
              Block("Built things.", "line 3", list_item=True)]
    m = build(Extracted("text", blocks))
    assert any("run backwards" in w for w in m.warnings)
