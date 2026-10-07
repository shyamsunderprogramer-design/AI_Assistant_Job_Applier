"""Resume Tailoring end to end, its page and routes, and the rest of the app still working."""

import io
import json

import pytest
from docx import Document

from ml.tailoring import documents, pipeline

JD = """Platform Engineer
Requirements:
- 3+ years of experience with Kubernetes and Terraform.
- Experience with AWS and CI/CD pipelines.
- Experience with Rust services.
Nice to have:
- Datadog dashboards experience.
"""


def sample_resume(path):
    d = Document()
    d.add_paragraph("Jordan Example")
    d.add_paragraph("jordan@example.com | +1 555 010 0199")
    d.add_heading("Experience", level=1)
    d.add_paragraph("Operations Analyst, Acme Corp     Jan 2019 – Dec 2023")
    d.add_paragraph("Ran Kubernetes 1.27 clusters on EKS for 40 services.", style="List Bullet")
    d.add_paragraph("Wrote Terraform modules for 12 AWS accounts.", style="List Bullet")
    d.add_paragraph("Built CI/CD pipelines in GitLab CI for every release.", style="List Bullet")
    d.add_heading("Skills", level=1)
    d.add_paragraph("Datadog, Python, Bash")
    d.add_heading("Education", level=1)
    d.add_paragraph("B.S. Computer Science, State University     2014 – 2018")
    d.save(path)
    return path


@pytest.fixture
def no_word(monkeypatch):
    monkeypatch.setattr(documents, "word_to_pdf", lambda docx, out: False)


@pytest.fixture
def no_network(monkeypatch):
    monkeypatch.setattr("ml.tailoring.versions._fetch", lambda product: None)


def test_a_whole_run_without_a_model(tmp_path, no_word, no_network, monkeypatch):
    # The draft step's rewrite needs a model; without one, the resume's own words are used.
    def no_model(*a, **k):
        raise RuntimeError("no model in tests")
    monkeypatch.setattr("ml.resume.tailor.complete", no_model)
    resume = sample_resume(tmp_path / "resume.docx")
    before = resume.read_bytes()
    stages = []
    (tmp_path / "run").mkdir()
    result = pipeline.run(tmp_path / "run", resume, JD, None, None,
                          progress=lambda s, note="": stages.append(s), use_model=False)
    assert list(dict.fromkeys(stages)) == list(pipeline.STAGES)          # in order; a stage may report progress
    assert resume.read_bytes() == before                                   # the upload is never changed
    statuses = {m["requirement"]["normalized"]: m["status"] for m in result["report"]["matches"]}
    assert statuses["kubernetes"] == "direct" and statuses["rust"] == "not_found"
    assert result["files"] == {"docx": "tailored-resume.docx", "pdf": "tailored-resume.pdf"}
    assert result["pages"]["pdf"]                                          # a preview exists
    assert any(g["requirement"]["normalized"] == "rust" for g in result["gaps"])
    assert any(v["tool"].lower() == "kubernetes" for v in result["versions"])  # checked, Unknown offline


def test_the_docx_opens_and_the_pdf_matches_it(tmp_path, no_word, no_network, monkeypatch):
    monkeypatch.setattr("ml.resume.tailor.complete", lambda *a, **k: (_ for _ in ()).throw(RuntimeError()))
    (tmp_path / "run").mkdir()
    result = pipeline.run(tmp_path / "run", sample_resume(tmp_path / "r.docx"), JD, None, None, use_model=False)
    docx = Document(str(tmp_path / "run" / result["files"]["docx"]))
    paragraphs = [p.text for p in docx.paragraphs if p.text.strip()]
    from ml.tailoring.documents import _lines
    pdf_words = set(" ".join(r[6] for r in _lines(tmp_path / "run" / result["files"]["pdf"])).lower().split())
    for text in paragraphs:
        words = [w for w in text.lower().split() if w.isalpha()]
        assert sum(w in pdf_words for w in words) >= 0.9 * len(words), text
    assert not [c for c in result["checks"] if c["severity"] == "problem"], result["checks"]


# -- the page and its routes -----------------------------------------------------------

@pytest.fixture
def client(monkeypatch, tmp_path):
    from backend.api import app as webapp
    monkeypatch.setattr(pipeline, "RUNS", tmp_path / "runs")
    started = []
    monkeypatch.setattr(pipeline, "start", lambda *a: started.append(a))
    monkeypatch.setattr(webapp, "cfg", lambda: None)
    c = webapp.app.test_client()
    c.started = started
    return c


def test_the_tailor_page_loads(client):
    page = client.get("/tailor")
    assert page.status_code == 200 and b"Create tailored resume" in page.data


@pytest.mark.parametrize("files,form,which,words", [
    ({}, {"jd_text": JD}, "resume", "Choose your resume"),
    ({"resume": (io.BytesIO(b"x"), "cv.png")}, {"jd_text": JD}, "resume", ".docx or .pdf"),
    ({"resume": (io.BytesIO(b"x"), "cv.docx")}, {"jd_text": ""}, "job", "Paste the job description"),
])
def test_missing_or_unsupported_files_are_named(client, files, form, which, words):
    reply = client.post("/tailor/start", data={**form, **files}, content_type="multipart/form-data")
    body = reply.get_json()
    assert reply.status_code == 400 and body["file"] == which and words in body["error"]
    assert client.started == []


def test_a_good_upload_starts_a_run(client):
    reply = client.post("/tailor/start", data={"jd_text": JD, "resume": (io.BytesIO(b"PK.."), "cv.docx")},
                        content_type="multipart/form-data")
    run = reply.get_json()["run"]
    assert reply.get_json()["ok"] and client.started
    assert client.get(f"/tailor/{run}/status").get_json()["ok"]
    assert client.get("/tailor/../../etc/status").status_code == 404     # never outside the runs


def test_the_rest_of_the_app_still_works():
    from backend.api import app as webapp
    from backend.config.loader import load_config
    from data_engineering.db import session
    if session._SessionFactory is None:
        session.init_engine(load_config().database_url)
    webapp.app.config["TESTING"] = True
    with webapp.app.test_client() as client:
        for path in ("/", "/setup", "/assistant", "/applicant"):
            assert client.get(path).status_code == 200, path
        assert b'href="/tailor"' in client.get("/").data          # reachable from the jobs page


def test_a_job_score_is_saved_and_reused(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "SCORES", tmp_path / "scores")
    resume = sample_resume(tmp_path / "r.docx")
    first = pipeline.job_score(7, resume, JD)
    assert 0 < first["score"] < 100 and "Rust" in first["gaps"]
    monkeypatch.setattr("ml.tailoring.matching.match", lambda *a, **k: 1 / 0)   # not recomputed
    assert pipeline.job_score(7, resume, JD) == first
    assert pipeline.job_score(8, resume, "too short")["score"] is None


def test_the_tailor_page_takes_a_job(client):
    assert client.get("/tailor?job=abc").status_code == 200


def test_a_run_from_a_job_page_uses_the_saved_resume(client, monkeypatch, tmp_path):
    saved = sample_resume(tmp_path / "base_resume.docx")
    monkeypatch.setattr("ml.resume.pipeline.base_resume_path", lambda cfg: saved)
    reply = client.post("/tailor/start", data={"jd_text": JD, "use_saved": "1"},
                        content_type="multipart/form-data")
    assert reply.get_json()["ok"], reply.get_json()
    folder, resume_path = client.started[0][0], client.started[0][1]
    assert resume_path.read_bytes() == saved.read_bytes() and resume_path.parent == folder / "inputs"


def test_a_job_run_puts_its_files_where_the_job_page_looks(tmp_path, no_word, no_network, monkeypatch):
    monkeypatch.setattr("ml.resume.tailor.complete", lambda *a, **k: (_ for _ in ()).throw(RuntimeError()))
    monkeypatch.setattr(pipeline, "JOBS", tmp_path / "jobs")
    folder = tmp_path / "run"; folder.mkdir()
    target = tmp_path / "output" / "acme_platform-engineer.docx"
    pipeline.remember_job_run(42, folder, target)
    pipeline.run(folder, sample_resume(tmp_path / "r.docx"), JD, None, None, use_model=False)
    pipeline.place_for_job(folder, target)
    assert target.exists() and target.with_suffix(".pdf").exists()
    assert pipeline.job_run(42) == "run" and pipeline.job_of_run(folder)["job_id"] == 42


def test_a_lost_term_is_won_back_by_refitting_before_another_model_round(monkeypatch):
    """A draft short of the goal is refitted (no model) first; only if that fails is it rewritten."""
    from ml.tailoring import rounds
    made, refits = [], []

    class D:
        accepted = True

        def __init__(self, after):
            self.after = after

    def make(extra, focus):
        made.append(focus)
        return D(91)

    def refit(draft, focus):
        refits.append(focus)
        return D(95)
    monkeypatch.setattr(rounds, "ats_of", lambda d, jd, company=None:
                        {"before": 94, "after": d.after, "fixable": ["platform"] if d.after < 94 else [],
                         "concepts": []})
    best, ats, used = rounds.best_draft(make, "jd", refit=refit)
    assert ats["after"] == 95 and used == 1 and len(made) == 1 and refits == [["platform"]]


def test_a_rewrite_that_loses_a_required_items_only_proof_is_undone():
    from types import SimpleNamespace as NS

    from ml.resume.tailor import TailorResult
    from ml.tailoring.draft import keep_required_wording
    line = "Led incident response for P1 outages across 40 services, cutting MTTR by 35%."
    result = TailorResult(summary=None)
    result.bullets = [{"original": "- " + line, "tailored": "Restored P1 outages across 40 services, cutting MTTR 35%."},
                      {"original": "Built Terraform modules.", "tailored": "Built Terraform modules for AWS."}]
    report = NS(matches=[NS(status="direct", evidence=[{"text": line}],
                            requirement={"priority": "required", "normalized": "incident response"}),
                         NS(status="direct", evidence=[{"text": "Built Terraform modules."}],
                            requirement={"priority": "required", "normalized": "terraform"})])
    assert keep_required_wording(result, report) == [line]
    assert [b["original"] for b in result.bullets] == ["Built Terraform modules."]   # still says Terraform: kept
