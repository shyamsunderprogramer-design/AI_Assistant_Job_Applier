"""Every template offered for applying is read back whole and in order by every reader an ATS uses."""

import pytest
from docx import Document

from ml.resume.ats_readback import Reading, _check, verify_template
from ml.resume.parser import parse_resume
from ml.resume.style import TEMPLATES, for_job

# A made-up person. Short degree names with their dates far off to the right, and bullets of long
# words that justified lines spread wide: the two things pdfminer misread in Oct 2026.
LINES = [
    "Jordan Example",
    "Denver, CO · jordan@example.com · +1 555 010 0199",
    "Summary",
    "Platform engineer with nine years of experience building delivery pipelines and infrastructure "
    "automation for regulated workloads across three clouds.",
    "Skills",
    "Cloud Platforms: AWS (EC2, S3, Lambda), Azure (AKS, Key Vault), GCP (GKE, Cloud Build)",
    "Automation: Terraform, Ansible, Helm, Python, Bash",
    "Experience",
    "SENIOR PLATFORM ENGINEER, ACME CORP, USA\tFeb 2021 – Present",
    "• Orchestrated quarterly infrastructure modernization, consolidating Kubernetes environments and "
    "networking components, improving deployment throughput by 25% and eliminating bottlenecks for "
    "customer-facing applications.",
    "• Standardized infrastructure verification using Checkov, tfsec, and Conftest, intercepting misconfigurations "
    "before deployment across forty repositories.",
    "PLATFORM ENGINEER, INITECH, USA\tJun 2016 – Jan 2021",
    "• Documented troubleshooting procedures, synthetic monitoring, and automatic remediation, shortening "
    "recovery time by 35%.",
    "Education",
    "Master of Science\tJan 2014 – Dec 2015",
    "Example State University, Denver, CO, USA",
    "Bachelor of Technology\tJul 2009 – Apr 2013",
    "Example Institute of Technology, Example City, India",
]


@pytest.fixture
def base(tmp_path):
    src = tmp_path / "master.docx"
    doc = Document()
    for line in LINES:
        doc.add_paragraph(line)
    doc.save(src)
    return parse_resume(src)


@pytest.mark.parametrize("key", [k for k, st in TEMPLATES.items() if st.ats_safe])
def test_every_template_offered_for_applying_reads_back_whole(base, tmp_path, key):
    proof = verify_template(key, base, None, tmp_path)
    assert proof["ok"], proof["readers"]


def test_the_sidebar_is_not_offered_for_applying():
    assert TEMPLATES["sidebar"].ats_safe is False
    assert for_job(default="sidebar")[0] == "classic"


def test_a_degree_and_its_dates_are_one_line_to_pdfminer(base, tmp_path):
    from pdfminer.high_level import extract_text
    verify_template("classic", base, None, tmp_path)
    lines = [l.strip() for l in extract_text(str(tmp_path / "classic.pdf")).splitlines()]
    master = next(l for l in lines if l.startswith("Master of Science"))
    assert "Jan 2014" in master and "Dec 2015" in master


@pytest.mark.parametrize("slack", [8, 30, 60])
def test_a_widely_spaced_justified_line_reads_in_one_piece(tmp_path, slack):
    """A justified line with few words and much room: pdfminer read it as pieces (Oct 2026)."""
    from pdfminer.high_level import extract_text
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from ml.resume import pdf as resume_pdf
    font = resume_pdf.FONT
    first, long_word = "and load it is li if il", "Infrastructure-automation"
    width = stringWidth(first, font, 11) + slack
    assert stringWidth(" " + long_word, font, 11) > slack          # the long word goes to the next line
    style = ParagraphStyle("b", fontName=font, fontSize=11, leading=13, alignment=TA_JUSTIFY, splitLongWords=0)
    out = tmp_path / "line.pdf"
    # The frame pads its width by 6 points a side.
    SimpleDocTemplate(str(out), pagesize=(612, 792), leftMargin=72, rightMargin=612 - 72 - width - 12).build(
        [Paragraph(f"{first} {long_word} applications, read whole.", style)])
    lines = [" ".join(l.split()) for l in extract_text(str(out)).splitlines() if l.strip()]
    assert lines[0] == first, lines


def test_the_federal_template_writes_months_as_numbers(base, tmp_path):
    from pypdf import PdfReader
    verify_template("federal", base, None, tmp_path)
    pdf = "\n".join(p.extract_text() for p in PdfReader(str(tmp_path / "federal.pdf")).pages)
    docx = "\n".join(p.text for p in Document(str(tmp_path / "federal.docx")).paragraphs)
    for text in (pdf, docx):
        assert "02/2021 – Present" in text and "06/2016 – 01/2021" in text and "01/2014 – 12/2015" in text
        assert "Feb 2021" not in text


def test_numeric_dates_change_only_the_dates():
    from ml.resume.writer import numeric_dates
    assert numeric_dates("LEAD ENGINEER, MAY CORP\tSept. 2019 – Mar 2021") == "LEAD ENGINEER, MAY CORP\t09/2019 – 03/2021"
    assert numeric_dates("Engineer | 2019 – 2021") == "Engineer | 2019 – 2021"
    assert numeric_dates("Built things in March 2020") == "Built things in March 2020"   # no dates column


def test_the_check_catches_dates_read_apart_and_a_heading_run_into_its_text():
    blocks = [("name", "Jordan Example"), ("heading", "Education"),
              ("degree", "Master of Science\tJan 2014 – Dec 2015"), ("school", "Example State University"),
              ("degree", "Bachelor of Technology\tJul 2009 – Apr 2013"), ("school", "Example Institute of Technology")]
    apart = Reading("pdfminer", "Jordan Example\nEDUCATION Master of Science\nExample State University\n"
                                "Bachelor of Technology\n" + "Example Institute of Technology, " * 4 +
                                "\nJan 2014 – Dec 2015\nJul 2009 – Apr 2013\n")
    _check(apart, blocks)
    kinds = {f.kind for f in apart.findings}
    assert "heading" in kinds and "dates" in kinds
    whole = Reading("pypdf", "Jordan Example\nEDUCATION\nMaster of Science  Jan 2014 – Dec 2015\n"
                             "Example State University\nBachelor of Technology  Jul 2009 – Apr 2013\n"
                             "Example Institute of Technology\n")
    _check(whole, blocks)
    assert whole.findings == []


def test_a_clean_read_back_is_one_note(base, tmp_path):
    from ml.tailoring.documents import Rendered, ats_checks
    from ml.resume.tailor import TailorResult
    from ml.resume.writer import plan_document
    verify_template("modern", base, None, tmp_path)
    rendered = Rendered(tmp_path / "modern.docx", tmp_path / "modern.pdf", None)
    found = ats_checks(rendered, plan_document(base, TailorResult(summary=None, template="modern")))
    assert [c.severity for c in found] == ["note"] and "whole and in order" in found[0].detail


@pytest.mark.parametrize("source, company, text, clearance, want", [
    ("usajobs", "Department of Veterans Affairs", "IT Specialist", "", "federal"),
    ("greenhouse", "U.S. Department of State", "Cloud engineer", "", "federal"),
    ("greenhouse", "Example Federal Partners", "Must hold an active Secret clearance.", "", "federal"),
    ("lever", "Example Systems", "Cloud engineer", "Top Secret", "federal"),
    ("lever", "Example Systems", "This position requires Public Trust.", "", "federal"),
    ("greenhouse", "U.S. Bank", "Cloud engineer", "", "classic"),
    ("greenhouse", "Example Security", "Our platform serves federal customers and holds FedRAMP.", "", "classic"),
    ("ashby", "Example Labs", "We are a Series B startup building developer tools.", "", "modern"),
    ("ashby", "Example Labs", "We move fast like a startup.", "", "classic"),
    ("ashby", "Example Labs", "Kubernetes and Terraform.", "", "classic"),
])
def test_the_template_for_a_job(source, company, text, clearance, want):
    template, why = for_job(source, company, "Engineer", text, default="classic", clearance=clearance)
    assert template == want and why


def test_the_persons_default_stands_unless_the_job_is_federal():
    assert for_job("ashby", "Example Labs", "Engineer", "We are a Series B startup.", default="compact")[0] == "compact"
    assert for_job("usajobs", "", "Engineer", "", default="compact")[0] == "federal"
