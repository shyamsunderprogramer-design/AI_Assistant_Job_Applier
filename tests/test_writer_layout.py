"""How the generated .docx is laid out, and what must never regress.

Every rule here is a defect measured on a real output of the first writer:

    paragraph styles used     : {'Normal': 145}   -- no headings at all
    bullets with hanging indent:   0 of 43
    long body lines, no bullet :  64

Four bullets sat under one heading with a marker on exactly one of them,
because bullets were decided by whether the SOURCE line still carried its
marker -- and PDF extraction drops that marker from some lines and not
others. Wrapped bullet text sat under the dot instead of under the text.
"""

from pathlib import Path

import pytest
from docx import Document

from resume.parser import Resume, Section
from resume.tailor import TailorResult
from resume.writer import write_tailored_resume

BULLET = "•"


@pytest.fixture
def written(tmp_path):
    """A resume whose experience bullets have lost most of their markers."""
    base = Resume(raw_text="", sections=[
        Section(heading="HEADER", lines=[
            "Jane Roe",
            "Boston, MA | +1 555 0100 | jane@example.com",
        ]),
        Section(heading="EXPERIENCE", lines=[
            "SENIOR PLATFORM ENGINEER, ACME CORP, USA\tFeb 2025 - Present",
            "Reliability & Scale",
            "Ran the release train for eleven services across three regions every week.",
            "Cut the median deploy from forty minutes to under six by reworking the pipeline.",
            f"{BULLET} Held the on-call rota and wrote the runbooks the rota depends on.",
            "Built the golden-path templates every new service starts from today.",
        ]),
        Section(heading="EDUCATION", lines=[
            "Master of Science\tJan 2023 - Dec 2024",
        ]),
    ])
    result = TailorResult(summary="Platform engineer, eleven years.", bullets=[])
    path = write_tailored_resume(base, result, tmp_path / "out.docx")
    return Document(path)


def body(doc):
    return [p for p in doc.paragraphs if p.text.strip()]


def bullets(doc):
    return [p for p in body(doc) if p.text.strip().startswith(BULLET)]


# --- the bug that started this -------------------------------------------

def test_a_section_bullets_all_of_its_body_or_none_of_it(written):
    """One marker survived extraction out of four. All four are bullets."""
    lines = [p.text for p in bullets(written)]
    assert any("release train" in t for t in lines)
    assert any("median deploy" in t for t in lines)
    assert any("on-call rota" in t for t in lines)
    assert any("golden-path" in t for t in lines)


def test_every_bullet_hangs(written):
    """Without a negative first-line indent, wrapped text sits under the dot."""
    marks = bullets(written)
    assert marks, "no bullets were written at all"
    for para in marks:
        assert (para.paragraph_format.first_line_indent or 0) < 0, para.text[:40]


def test_sections_are_real_headings_not_bold_normal_text(written):
    """Parsers locate EXPERIENCE and EDUCATION by style, not by reading."""
    headings = [p.text.strip() for p in body(written)
                if p.style.name.startswith("Heading")]
    assert "EXPERIENCE" in headings
    assert "EDUCATION" in headings


def test_a_heading_is_never_stranded_at_the_foot_of_a_page(written):
    for para in body(written):
        if para.style.name.startswith("Heading"):
            assert para.paragraph_format.keep_with_next


# --- layout -------------------------------------------------------------

def test_a_job_keeps_its_dates_on_the_right(written):
    role = next(p for p in body(written) if "ACME CORP" in p.text)
    assert "\t" in role.text
    assert len(role.paragraph_format.tab_stops) == 1
    assert role.paragraph_format.keep_with_next, "a job must not split from its bullets"


def test_a_job_title_is_not_turned_into_a_bullet(written):
    for para in bullets(written):
        assert "ACME CORP" not in para.text
        assert "Master of Science" not in para.text


def test_a_grouping_label_is_not_turned_into_a_bullet(written):
    """"Reliability & Scale" labels the bullets under it; it is not one."""
    label = next(p for p in body(written) if p.text.strip() == "Reliability & Scale")
    assert not label.text.startswith(BULLET)
    assert label.runs[0].bold


def test_the_contact_line_uses_one_separator(written):
    contact = body(written)[1]
    assert "|" not in contact.text
    assert "·" in contact.text


def test_the_name_leads(written):
    name = body(written)[0]
    assert name.text.strip() == "Jane Roe"
    assert name.runs[0].bold


# --- what ATS parsers choke on, and must stay absent ---------------------

def test_nothing_a_parser_chokes_on_is_emitted(written):
    assert written.tables == []
    for section in written.sections:
        assert not [p for p in section.header.paragraphs if p.text.strip()]
        assert not [p for p in section.footer.paragraphs if p.text.strip()]
    assert not written.inline_shapes


# --- omission: the half of tailoring that was never wired up --------------

def _plan(omitted, lines=None):
    from resume.writer import plan_document

    base = Resume(raw_text="", sections=[
        Section(heading="HEADER", lines=["Jane Roe", "jane@example.com"]),
        Section(heading="EXPERIENCE", lines=lines or [
            "SENIOR PLATFORM ENGINEER, ACME CORP\tFeb 2025 - Present",
            "Ran the weekly release train for eleven services across three regions.",
            "Cut the median deploy from forty minutes to under six by reworking it all.",
            "Maintained the legacy Perl reporting jobs nobody else would touch here.",
            "Held the on-call rota and wrote the runbooks that the rota depends on.",
        ]),
    ])
    return plan_document(base, TailorResult(summary=None, bullets=[], omitted=omitted))


def _texts(plan):
    return [t for _, t in plan]


def test_a_named_bullet_is_actually_removed():
    """The model reported omissions and the writer emitted everything anyway,
    so every tailored resume came out the same length as the original."""
    plan = _plan([{"original": "Maintained the legacy Perl reporting jobs nobody"
                               " else would touch here.",
                   "reason": "irrelevant to this role"}])
    assert not any("Perl" in t for t in _texts(plan))
    assert any("release train" in t for t in _texts(plan))


def test_a_prose_omission_removes_nothing():
    """The old shape: a description names no line, so nothing can be dropped."""
    plan = _plan(["The legacy Perl work was omitted as irrelevant."])
    assert any("Perl" in t for t in _texts(plan))


def test_a_job_title_is_never_dropped():
    plan = _plan([{"original": "SENIOR PLATFORM ENGINEER, ACME CORP\tFeb 2025 - Present",
                   "reason": "not relevant"}])
    assert any("ACME CORP" in t for t in _texts(plan))


def test_a_section_never_gives_up_more_than_half():
    """A role stripped to one line reads as a gap in the career."""
    body = [
        "Ran the weekly release train for eleven services across three regions.",
        "Cut the median deploy from forty minutes to under six by reworking it all.",
        "Maintained the legacy Perl reporting jobs nobody else would touch here.",
        "Held the on-call rota and wrote the runbooks that the rota depends on.",
    ]
    plan = _plan([{"original": b, "reason": "x"} for b in body],
                 lines=["SENIOR PLATFORM ENGINEER, ACME CORP\tFeb 2025 - Present", *body])
    kept = [t for k, t in plan if k == "bullet"]
    assert len(kept) >= 2, "the whole job was gutted"
